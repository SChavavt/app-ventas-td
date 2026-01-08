#app_admin.py
import streamlit as st
import streamlit.components.v1 as components
import json
import time
import random
import html
import re
import pandas as pd
import boto3
from botocore.exceptions import ClientError
import gspread
from gspread.exceptions import APIError
from gspread.utils import rowcol_to_a1
from google.oauth2.service_account import Credentials as GoogleCredentials
from io import BytesIO
from datetime import datetime, date
from zoneinfo import ZoneInfo
import os
import uuid
from urllib.parse import urlparse, unquote
from contextlib import suppress
from streamlit.runtime.scriptrunner import StopException
import numbers

# Reintentos robustos para Google Sheets
RETRIABLE_CODES = {429, 500, 502, 503, 504}
TRANSIENT_TEXT_MARKERS = {
    "ratelimit",
    "ratelimitexceeded",
    "rate_limit_exceeded",
    "user_rate_limit_exceeded",
    "quotaexceeded",
    "quota_exceeded",
    "limitexceeded",
    "resource_exhausted",
    "backenderror",
}

REFRESH_COOLDOWN = 60
QUOTA_ERROR_THRESHOLD = 5


MOTIVO_RECHAZO_CANCELACION_COL = "Motivo_Rechazo/Cancelacion"
FECHA_CONFIRMADO_COL = "Fecha_Confirmado"
ESTADO_ENTREGA_COL = "Estado_Entrega"
GUIAS_DEVOLUCION_COL = "Numero_Guias_Devolucion"
GUIAS_DEVOLUCION_OPTIONS = ["0 guías", "1 guía", "2 guías"]
CDMX_TIMEZONE = ZoneInfo("America/Mexico_City")


CONFIRMADOS_SYNC_COLUMN_MAP = {
    "Estado_Surtido_Almacen": "Estado_Surtido_Almacen",
    "Fecha_Entrega": "Fecha_Entrega",
    FECHA_CONFIRMADO_COL: FECHA_CONFIRMADO_COL,
    ESTADO_ENTREGA_COL: ESTADO_ENTREGA_COL,
}


COLUMNAS_OBJETIVO_CONFIRMADOS = [
    "ID_Pedido",
    "Hora_Registro",
    "Folio_Factura",
    "Folio_Factura_Refacturada",
    "Cliente",
    "Vendedor_Registro",
    "Tipo_Envio",
    "Fecha_Entrega",
    "Estado_Surtido_Almacen",
    ESTADO_ENTREGA_COL,
    "Estado_Pago",
    "Comprobante_Confirmado",
    FECHA_CONFIRMADO_COL,
    "Refacturacion_Tipo",
    "Refacturacion_Subtipo",
    "Forma_Pago_Comprobante",
    "Monto_Comprobante",
    "Fecha_Pago_Comprobante",
    "Banco_Destino_Pago",
    "Terminal",
    "Referencia_Comprobante",
    "Link_Adjuntos",
    "Modificacion_Surtido",
    "Link_Adjuntos_Modificacion",
    "Link_Refacturacion",
    "Link_Adjuntos_Guia",
    "Motivo_NotaVenta",
    MOTIVO_RECHAZO_CANCELACION_COL,
]


ESTADO_ENTREGA_OPCIONES = ["📦 Entregado", "⏳ No Entregado"]
ESTADO_ENTREGA_DEFAULT = ESTADO_ENTREGA_OPCIONES[1]


def obtener_fecha_confirmado_cdmx() -> str:
    """Devuelve la fecha-hora actual en CDMX con formato legible para la hoja."""
    return datetime.now(CDMX_TIMEZONE).strftime("%Y-%m-%d %H:%M:%S")


def normalize_estado_entrega(value) -> str:
    """Normaliza valores del estado de entrega al formato con emoji."""
    raw = str(value or "").strip()
    if not raw or raw.lower() == "nan":
        return ""

    raw_lower = raw.lower()
    for opcion in ESTADO_ENTREGA_OPCIONES:
        if raw_lower == opcion.lower():
            return opcion

    if "entregado" in raw_lower and "no" not in raw_lower:
        return ESTADO_ENTREGA_OPCIONES[0]
    if "no entreg" in raw_lower:
        return ESTADO_ENTREGA_OPCIONES[1]

    return raw


def normalize_user_field(value: str | None) -> str:
    """Normaliza campos de usuario para mostrarlos solo si traen información."""
    raw = str(value or "").strip()
    if not raw or raw.lower() in {"nan", "none"}:
        return ""

    return raw


def ensure_sheet_column(worksheet, headers: list[str], column_name: str) -> list[str]:
    """Garantiza que exista una columna en la hoja de Google Sheets."""
    headers_list = list(headers) if headers else []
    if column_name in headers_list:
        return headers_list

    desired_index = len(headers_list) + 1

    try:
        col_count = getattr(worksheet, "col_count", desired_index - 1) or 0
    except Exception:
        col_count = desired_index - 1

    if col_count < desired_index:
        with suppress(Exception):
            worksheet.add_cols(desired_index - col_count)

    try:
        safe_batch_update(
            worksheet,
            [
                {
                    "range": rowcol_to_a1(1, desired_index),
                    "values": [[column_name]],
                }
            ],
        )
        headers_list.append(column_name)
    except Exception as err:
        st.warning(
            f"⚠️ No se pudo asegurar la columna '{column_name}' en la hoja: {err}"
        )

    return headers_list


def _filter_cancelled_pedidos(df: pd.DataFrame) -> pd.DataFrame:
    """Oculta pedidos cancelados marcados en la columna de motivos."""
    if df is None or df.empty:
        return df

    trabajo = df.copy()
    if MOTIVO_RECHAZO_CANCELACION_COL in trabajo.columns:
        mask_cancelados = (
            trabajo[MOTIVO_RECHAZO_CANCELACION_COL]
            .astype(str)
            .str.contains(r"Cancelado\[", na=False)
        )
        trabajo = trabajo[~mask_cancelados]
    return trabajo


def refresh_pedidos_pagados_no_confirmados(
    df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Recalcula y guarda en sesión los pedidos con comprobante pendiente."""

    if df is None:
        df = st.session_state.get("df_pedidos")

    if df is None or df.empty:
        pendientes = pd.DataFrame()
    else:
        trabajo = df.copy()

        if "Comprobante_Confirmado" in trabajo.columns:
            mask_no_confirmado = (
                trabajo["Comprobante_Confirmado"].astype(str).str.strip().str.lower()
                != "sí"
            )
            pendientes = trabajo[mask_no_confirmado].copy()
        else:
            pendientes = trabajo.iloc[0:0].copy()

        if "Tipo_Envio" in pendientes.columns:
            pendientes = pendientes[
                ~pendientes["Tipo_Envio"].isin(
                    ["🎓 Cursos y Eventos", "📋 Solicitudes de Guía"]
                )
            ].copy()

    if not pendientes.empty:
        pendientes = _filter_cancelled_pedidos(pendientes)
        pendientes = pendientes.drop(columns=["display_label"], errors="ignore")

    st.session_state.pedidos_pagados_no_confirmados = pendientes
    return pendientes


if "pedidos_reload_nonce" not in st.session_state:
    st.session_state["pedidos_reload_nonce"] = 0

if "comprobante_form_nonce" not in st.session_state:
    st.session_state["comprobante_form_nonce"] = 0


def allow_refresh(key: str, container=st, cooldown: int = REFRESH_COOLDOWN) -> bool:
    """Rate-limit manual reloads to avoid hitting Google Sheets too often."""
    now = time.time()
    last = st.session_state.get(key)
    if last and now - last < cooldown:
        container.warning("⚠️ Se recargó recientemente. Espera unos segundos.")
        return False
    st.session_state[key] = now
    return True

def _err_signature(e) -> tuple[int|None, str]:
    """Extrae status y texto para decidir si reintentar."""
    status = getattr(getattr(e, "response", None), "status_code", None)
    try:
        text = e.response.text  # puede incluir 'rateLimitExceeded', 'USER_RATE_LIMIT_EXCEEDED', etc.
    except Exception:
        text = str(e)
    return status, text.lower()


def _is_transient_quota_error(status: int | None, text: str) -> bool:
    """Determina si el error es transitorio y conviene reintentar."""
    if status in RETRIABLE_CODES:
        return True
    return any(marker in text for marker in TRANSIENT_TEXT_MARKERS)


def _register_quota_hit() -> int:
    """Incrementa el contador de quota y devuelve el total acumulado."""
    hits = st.session_state.get("_quota_hits", 0) + 1
    st.session_state["_quota_hits"] = hits
    return hits


def expand_link_adjuntos_columns(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Expande la columna Link_Adjuntos en columnas individuales ordenadas y únicas."""
    df_expanded = df.copy()

    columnas_expandidas_existentes = [
        c for c in df_expanded.columns if re.fullmatch(r"Link_Adjuntos_\d+", c)
    ]
    if columnas_expandidas_existentes:
        df_expanded = df_expanded.drop(columns=columnas_expandidas_existentes)

    if "Link_Adjuntos" not in df_expanded.columns:
        return df_expanded, []

    def _split_links(valor: str) -> list[str]:
        bruto = str(valor or "")
        partes = [p.strip() for p in bruto.replace("\n", ",").split(",") if p and p.strip()]
        return list(dict.fromkeys(partes))

    enlaces_por_fila = df_expanded["Link_Adjuntos"].fillna("").apply(_split_links)
    max_enlaces = int(enlaces_por_fila.map(len).max() or 0) if not enlaces_por_fila.empty else 0

    patron_columna_adjuntos = re.compile(r"Link_Adjuntos_\d+")

    columnas_creadas: list[str] = []
    for idx in range(max_enlaces):
        nombre_columna = f"Link_Adjuntos_{idx + 1}"
        columnas_creadas.append(nombre_columna)
        df_expanded[nombre_columna] = enlaces_por_fila.apply(
            lambda enlaces, i=idx: enlaces[i] if len(enlaces) > i else ""
        )

    columnas_creadas = [col for col in columnas_creadas if patron_columna_adjuntos.fullmatch(col)]

    return df_expanded, columnas_creadas

def safe_open_worksheet(sheet_id: str, worksheet_name: str, retries: int = 3):
    """
    Abre una worksheet con reintentos automáticos en caso de errores temporales
    utilizando backoff exponencial con jitter. Reutiliza la instancia cacheada
    del spreadsheet y evita limpiar recursos globales.
    """
    now = time.time()
    locked_until = st.session_state.get("_quota_locked_until")
    if locked_until is not None:
        if now >= locked_until:
            st.session_state["_quota_hits"] = 0
            st.session_state.pop("_quota_locked_until", None)
        else:
            st.error("🚫 Se alcanzó el límite de cuota de Google Sheets. Espera antes de reintentar.")
            raise RuntimeError("google-sheets quota exceeded")

    if st.session_state.get("_quota_hits", 0) >= QUOTA_ERROR_THRESHOLD:
        st.session_state["_quota_locked_until"] = now + REFRESH_COOLDOWN
        st.error("🚫 Se alcanzó el límite de cuota de Google Sheets. Espera antes de reintentar.")
        raise RuntimeError("google-sheets quota exceeded")

    last_err = None
    delay = 1.0
    ss = None
    for attempt in range(retries):
        try:
            if ss is None:
                ss = get_spreadsheet(sheet_id)  # usa instancia cacheada del spreadsheet
            ws = ss.worksheet(worksheet_name)
            st.session_state["_quota_hits"] = 0
            st.session_state.pop("_quota_locked_until", None)
            return ws
        except gspread.exceptions.APIError as e:
            last_err = e
            status, text = _err_signature(e)
            is_rate = _is_transient_quota_error(status, text)
            if is_rate:
                hits = _register_quota_hit()
                if hits >= QUOTA_ERROR_THRESHOLD:
                    st.session_state["_quota_locked_until"] = time.time() + REFRESH_COOLDOWN
                    st.error("🚫 Se detectaron múltiples errores de cuota. Espera antes de reintentar.")
                    break
            if is_rate and attempt < retries - 1:
                jitter = random.uniform(0, delay)
                st.warning(
                    f"⚠️ Error de Google Sheets al abrir '{worksheet_name}'. Reintentando en {delay + jitter:.1f}s..."
                )
                time.sleep(delay + jitter)
                delay *= 2
                continue
            break
    st.error(
        f"❌ No se pudo abrir la hoja '{worksheet_name}' tras {retries} intentos: {last_err}"
    )
    raise last_err


st.set_page_config(page_title="App Admin TD", layout="wide")


TAB_QUERY_PARAM = "tab"
TAB_SESSION_KEY = "admin_current_tab_index"


def _normalize_tab_index(raw_index: int | str | None, total: int) -> int:
    """Convierte el índice de pestaña en un entero válido dentro de rango."""
    if raw_index is None:
        candidate = 0
    else:
        try:
            candidate = int(raw_index)
        except (TypeError, ValueError):
            candidate = 0

    if total <= 0:
        return 0
    return max(0, min(total - 1, candidate))


def rerun_current_tab():
    """Rerun Streamlit keeping the current tab in query params."""
    current_index = st.session_state.get(TAB_SESSION_KEY)
    if current_index is None:
        current_index = 0
    st.session_state["current_tab"] = str(current_index)
    st.query_params[TAB_QUERY_PARAM] = str(current_index)
    st.rerun()


def _comprobante_form_key(base: str) -> str:
    nonce = st.session_state.get("comprobante_form_nonce", 0)
    return f"{base}__{nonce}"


def clear_comprobante_form_state():
    """Limpia los campos persistentes del formulario de comprobantes."""
    keys_to_clear = {
        "pago_doble_admin",
        "comprobante_local_no_pagado",
        "fecha_pago_local",
        "forma_pago_local",
        "monto_pago_local",
        "terminal_local",
        "banco_destino_local",
        "referencia_local",
        "cp_pago1_admin",
        "fecha_pago1_admin",
        "forma_pago1_admin",
        "monto_pago1_admin",
        "terminal1_admin",
        "banco1_admin",
        "ref1_admin",
        "cp_pago2_admin",
        "fecha_pago2_admin",
        "forma_pago2_admin",
        "monto_pago2_admin",
        "terminal2_admin",
        "banco2_admin",
        "ref2_admin",
        "estado_entrega_local",
    }

    current_nonce = st.session_state.get("comprobante_form_nonce", 0)

    for key in keys_to_clear:
        st.session_state.pop(key, None)
        st.session_state.pop(f"{key}__{current_nonce}", None)

    dynamic_prefixes = (
        "fecha_pago_",
        "forma_pago_",
        "banco_pago_",
        "terminal_pago_",
        "monto_pago_",
        "ref_pago_",
        "motivo_rechazo__",
        "motivo_cancelacion__",
        "show_rechazo__",
        "show_cancel__",
    )

    for key in list(st.session_state.keys()):
        if any(key.startswith(prefix) for prefix in dynamic_prefixes):
            st.session_state.pop(key, None)
        elif any(key.startswith(f"{base}__") for base in keys_to_clear):
            st.session_state.pop(key, None)

    st.session_state["comprobante_form_nonce"] = current_nonce + 1


@st.cache_resource(ttl=60)
def _get_ws_datos():
    """Devuelve la worksheet 'datos_pedidos' con reintentos (usa safe_open_worksheet)."""
    return safe_open_worksheet(GOOGLE_SHEET_ID, "datos_pedidos")


def safe_batch_update(
    worksheet,
    data,
    retries: int = 5,
    base_delay: float = 2.0,
    max_delay: float = 64.0,
) -> None:
    """Realiza ``batch_update`` con reintentos exponenciales ante errores temporales."""
    if st.session_state.get("_quota_hits", 0) >= QUOTA_ERROR_THRESHOLD:
        st.error("🚫 Se alcanzó el límite de cuota de Google Sheets. Espera antes de reintentar.")
        raise RuntimeError("quota cooldown")

    last_err: APIError | None = None
    delay = base_delay
    for attempt in range(retries):
        try:
            worksheet.batch_update(data)
            st.session_state["_quota_hits"] = 0
            return
        except APIError as e:
            status, text = _err_signature(e)
            if not _is_transient_quota_error(status, text):
                raise

            last_err = e
            hits = _register_quota_hit()
            if hits >= QUOTA_ERROR_THRESHOLD:
                st.error("🚫 Se detectaron múltiples errores de cuota. Espera antes de reintentar.")
                break

            if attempt >= retries - 1:
                break

            capped_delay = min(delay, max_delay)
            jitter = random.uniform(0, capped_delay)
            sleep_for = capped_delay + jitter
            st.warning(
                f"⚠️ Error de Google Sheets al actualizar. Reintentando en {sleep_for:.1f}s..."
            )
            time.sleep(sleep_for)
            delay = min(delay * 2, max_delay)

    st.error("🚫 Google Sheets está aplicando un cooldown de cuota. Intenta nuevamente en unos minutos.")
    raise RuntimeError("quota cooldown") from last_err


def clean_modificacion_surtido(value) -> str:
    """Normaliza el texto de modificación de surtido evitando valores vacíos o nulos."""
    if value is None:
        return ""

    text = str(value).strip()
    if not text:
        return ""

    if text.lower() in {"nan", "none", "null", "n/a", "na"}:
        return ""

    return text


def normalize_id_pedido(value) -> str:
    """Normaliza el ID de pedido eliminando espacios y sufijos decimales espurios."""
    if value is None:
        return ""

    text = str(value).replace("\u00a0", " ").strip()
    if not text:
        return ""

    lowered = text.lower()
    if lowered in {"nan", "none", "null", "n/a", "na"}:
        return ""

    candidate = text.replace(",", "")
    if "." in candidate:
        integer_part, decimal_part = candidate.split(".", 1)
        if integer_part.isdigit() and decimal_part.strip("0") == "":
            candidate = integer_part

    if candidate.isdigit():
        # Normaliza 00123, 123.0, etc. → "123"
        return str(int(candidate))

    return text


def normalize_folio_factura(value) -> str:
    """Limpia el folio de factura removiendo espacios y marcadores nulos."""
    if value is None:
        return ""

    text = str(value).replace("\u00a0", " ").strip()
    if not text:
        return ""

    lowered = text.lower()
    if lowered in {"nan", "none", "null", "n/a", "na"}:
        return ""

    return text



# --- Helpers de Adjuntos --------------------------------------------------


def build_adjuntos_map_from_pedidos(
    df: pd.DataFrame | None,
) -> tuple[dict[str, object], dict[str, object], dict[str, object]]:
    """Regresa tres mapas ID_Pedido normalizado ➜ adjuntos por cada columna conocida.

    El contrato espera un ``pd.DataFrame`` con ``ID_Pedido`` y, opcionalmente, las
    columnas ``Adjuntos``, ``Adjuntos_Surtido`` y ``Adjuntos_Guia``. Cada mapa es
    independiente y se omite cuando la columna no existe o la celda está vacía.
    """

    empty_result: tuple[dict[str, object], dict[str, object], dict[str, object]] = ({}, {}, {})
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return empty_result

    if "ID_Pedido" not in df.columns:
        return empty_result

    ids_normalizados = df["ID_Pedido"].apply(normalize_id_pedido)

    def _build_for_column(column_name: str) -> dict[str, object]:
        column_map: dict[str, object] = {}
        if column_name not in df.columns:
            return column_map

        series = df[column_name]
        for row_idx in df.index:
            pedido_id = ids_normalizados.get(row_idx, "")
            if not pedido_id:
                continue

            try:
                raw_value = series.loc[row_idx]
            except KeyError:
                continue

            try:
                if pd.isna(raw_value):
                    continue
            except Exception:
                pass

            if isinstance(raw_value, str) and not raw_value.strip():
                continue

            column_map[pedido_id] = raw_value

        return column_map

    adjuntos_map = _build_for_column("Adjuntos")
    adjuntos_surtido_map = _build_for_column("Adjuntos_Surtido")
    adjuntos_guia_map = _build_for_column("Adjuntos_Guia")

    return adjuntos_map, adjuntos_surtido_map, adjuntos_guia_map


def extract_comprobante_urls_from_adjuntos(value) -> list[str]:
    """Extrae todas las URLs disponibles desde el campo Adjuntos."""

    url_pattern = re.compile(r"https?://[^\s,;]+", re.IGNORECASE)
    results: list[str] = []
    seen: set[str] = set()

    def _add(url: str | None) -> None:
        if not url:
            return
        url_text = str(url).strip()
        if not url_text:
            return
        lowered = url_text.lower()
        if lowered in {"nan", "none", "null"}:
            return
        if url_text not in seen:
            seen.add(url_text)
            results.append(url_text)

    def _extract_from_text(text: str) -> None:
        for match in url_pattern.findall(text):
            _add(match)

    def _process(obj) -> None:
        if obj is None:
            return

        if isinstance(obj, dict):
            for value in obj.values():
                _process(value)
            return

        if isinstance(obj, (list, tuple, set)):
            for item in obj:
                _process(item)
            return

        if isinstance(obj, bytes):
            try:
                decoded = obj.decode("utf-8", errors="ignore")
            except Exception:
                return
            _process(decoded)
            return

        if isinstance(obj, str):
            text_value = obj.strip()
            if not text_value:
                return
            try:
                parsed = json.loads(text_value)
            except Exception:
                _extract_from_text(text_value)
                return
            else:
                _process(parsed)
                return

        stringified = str(obj)
        if stringified and stringified.lower() not in {"nan", "none", "null"}:
            _extract_from_text(stringified)

    _process(value)
    return results


def clean_cell_text(value) -> str:
    """Normaliza valores hacia cadenas limpias para escritura en Google Sheets."""
    if value is None:
        return ""
    try:
        if pd.isna(value):  # type: ignore[arg-type]
            return ""
    except Exception:
        pass
    text = str(value)
    if text.strip().lower() in {"nan", "none", "null"}:
        return ""
    return text


def build_link_fallback_map(
    df: pd.DataFrame | None,
    columns: list[str],
) -> dict[str, dict[str, str]]:
    """Crea un mapa ID normalizado ➜ valores existentes para columnas de enlaces."""

    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return {}

    if "ID_Pedido" not in df.columns:
        return {}

    ids_normalizados = df["ID_Pedido"].apply(normalize_id_pedido)
    resultado: dict[str, dict[str, str]] = {}

    columnas_validas = [col for col in columns if col in df.columns]
    if not columnas_validas:
        return {}

    for row_idx in df.index:
        pedido_id_norm = ids_normalizados.get(row_idx, "")
        if not pedido_id_norm:
            continue
        fila = df.loc[row_idx]
        valores = {
            col: clean_cell_text(fila.get(col))  # type: ignore[index]
            for col in columnas_validas
        }
        if any(valores.values()):
            resultado[pedido_id_norm] = valores

    return resultado


def resolve_adjuntos_link(
    pedido_id: str,
    normalized_id: str,
    map_data: dict[str, object] | None,
    *,
    map_label: str,
    category: str,
    page_title: str,
    s3_client_instance,
) -> tuple[str, bool]:
    """Construye un enlace desde los adjuntos crudos, generando un índice cuando aplica."""

    if not map_data:
        return "", False

    raw_adjuntos = None
    if normalized_id:
        raw_adjuntos = map_data.get(normalized_id)
    if raw_adjuntos is None and pedido_id:
        raw_adjuntos = map_data.get(normalize_id_pedido(str(pedido_id)))

    if raw_adjuntos is None:
        return "", False

    try:
        urls = extract_comprobante_urls_from_adjuntos(raw_adjuntos)
    except Exception as parse_err:
        print(
            "[adjuntos_parser_error]",
            json.dumps(
                {
                    "pedido_id": normalized_id or pedido_id,
                    "tipo": map_label,
                    "error": str(parse_err),
                },
                ensure_ascii=False,
            ),
        )
        return "", False

    if not urls:
        return "", False

    if len(urls) == 1:
        return urls[0], True

    index_url = build_and_upload_comprobante_index_from_urls(
        normalized_id or pedido_id,
        urls,
        s3_client_instance,
        category=category,
        page_title=page_title,
    )
    return (index_url or "\n".join(urls)), True


# --- GOOGLE SHEETS CONFIGURATION ---
GOOGLE_SHEET_ID = '1aWkSelodaz0nWfQx7FZAysGnIYGQFJxAN7RO3YgCiZY'


@st.cache_data(ttl=300)
def cargar_pedidos_desde_google_sheet(sheet_id, worksheet_name, _nonce: int = 0):
    # 1) Intenta leer con reintentos usando el helper
    try:
        ws = safe_open_worksheet(sheet_id, worksheet_name)
        raw_values = ws.get_values()

        if not raw_values:
            headers = []
            df = pd.DataFrame()
        else:
            headers = ["" if h is None else str(h) for h in raw_values[0]]
            num_cols = len(headers)

            if num_cols == 0:
                df = pd.DataFrame()
            else:
                rows = []
                for raw_row in raw_values[1:]:
                    row = list(raw_row[:num_cols])
                    if len(row) < num_cols:
                        row.extend([""] * (num_cols - len(row)))
                    row = ["" if cell is None else cell for cell in row]
                    if all(str(cell).strip() == "" for cell in row):
                        continue
                    rows.append(row)

                df = pd.DataFrame(rows, columns=headers) if rows else pd.DataFrame(columns=headers)

        # 🔧 Normalización idéntica o equivalente a la tuya actual
        def _clean(s):
            return str(s).replace("\u00a0", " ").strip().replace("  ", " ").replace(" ", "_")

        if not df.empty or headers:
            df.columns = [_clean(c) for c in df.columns]

        alias = {
            "Folio de Factura": "Folio_Factura",
            "Folio_Factura_": "Folio_Factura",
            "ID_Pedido_": "ID_Pedido"
        }
        df = df.rename(columns=alias)

        # Asegura columnas clave
        for col in ["Folio_Factura", "ID_Pedido"]:
            if col not in df.columns:
                df[col] = ""

        # Limpieza de literales NA
        for col in ["Folio_Factura", "ID_Pedido"]:
            df[col] = (
                df[col]
                .astype(str)
                .str.replace("\u00a0", " ", regex=False)
                .str.strip()
            )
        NA_LITERALS = {"", "n/a", "na", "nan", "ninguno", "none"}
        df[["Folio_Factura", "ID_Pedido"]] = df[["Folio_Factura", "ID_Pedido"]].apply(
            lambda s: s.mask(s.str.lower().isin(NA_LITERALS))
        )
        df = df.dropna(subset=["Folio_Factura", "ID_Pedido"], how="all")

        if "ID_Pedido" in df.columns:
            df["ID_Pedido"] = df["ID_Pedido"].apply(normalize_id_pedido)

        # 2) Guarda snapshot “último bueno” por si falla luego
        st.session_state[f"_lastgood_{worksheet_name}"] = (df.copy(), list(headers))
        return df, headers

    except gspread.exceptions.APIError as e:
        # 3) Fallback: usa el último snapshot bueno si existe
        snap = st.session_state.get(f"_lastgood_{worksheet_name}")
        if snap:
            st.warning(f"♻️ Google Sheets dio un error temporal al leer '{worksheet_name}'. Mostrando el último dato bueno en caché.")
            return snap[0], snap[1]
        # 4) Si no hay snapshot, devuelve vacío pero sin matar la app
        st.error(f"❌ No se pudo leer '{worksheet_name}' (Google API). Intenta el botón de Recargar. Detalle: {e}")
        return pd.DataFrame(), []


@st.cache_resource
def get_google_sheets_client():
    """
    Crea el cliente de Google Sheets con google.oauth2 (no oauth2client).
    No abre ningún spreadsheet aquí para evitar errores 429 al crear el cliente.
    """
    max_retries = 3
    delay = 1
    for attempt in range(max_retries):
        try:
            credentials_json_str = st.secrets["google_credentials"]
            creds_dict = json.loads(credentials_json_str)
            # Normaliza el private_key con saltos de línea reales
            if "private_key" in creds_dict:
                creds_dict["private_key"] = creds_dict["private_key"].replace("\\n", "\n").strip()

            scopes = [
                "https://www.googleapis.com/auth/spreadsheets",
                "https://www.googleapis.com/auth/drive",  # necesario si agregas hojas, etc.
            ]
            creds = GoogleCredentials.from_service_account_info(creds_dict, scopes=scopes)
            client = gspread.authorize(creds)
            return client

        except Exception as e:
            if attempt < max_retries - 1:
                st.warning(
                    f"⚠️ Error al autenticar con Google Sheets ({e}). Reintentando en {delay}s..."
                )
                time.sleep(delay)
                delay *= 2
                st.cache_resource.clear()
            else:
                st.error(
                    f"❌ No se pudo autenticar con Google Sheets tras {max_retries} intentos: {e}"
                )
                class _FakeResp:
                    def __init__(self, text):
                        self.text = text

                raise gspread.exceptions.APIError(_FakeResp(str(e)))


@st.cache_resource
def get_spreadsheet(sheet_id: str):
    """Abre un spreadsheet por ``sheet_id`` reutilizando instancias cacheadas."""
    try:
        gc = get_google_sheets_client()
    except gspread.exceptions.APIError as e:
        snap = st.session_state.get("_last_spreadsheet")
        if snap:
            st.warning("♻️ No se pudo autenticar con Google Sheets. Usando snapshot en caché.")
            return snap
        st.error(f"❌ No se pudo autenticar con Google Sheets: {e}")
        return None

    try:
        ss = gc.open_by_key(sheet_id)
        st.session_state["_last_spreadsheet"] = ss
        return ss
    except gspread.exceptions.APIError as e:
        snap = st.session_state.get("_last_spreadsheet")
        if snap:
            st.warning("♻️ Error al abrir el spreadsheet. Usando snapshot en caché.")
            return snap
        raise


if "df_pedidos" not in st.session_state or "headers" not in st.session_state:
    df_pedidos, headers = cargar_pedidos_desde_google_sheet(
        GOOGLE_SHEET_ID, "datos_pedidos", st.session_state["pedidos_reload_nonce"]
    )
    # Excluir pedidos de cursos y eventos para que no aparezcan en ningún flujo
    if 'Tipo_Envio' in df_pedidos.columns:
        df_pedidos = df_pedidos[
            ~df_pedidos['Tipo_Envio'].isin(['🎓 Cursos y Eventos', '📋 Solicitudes de Guía'])
        ].copy()
    if MOTIVO_RECHAZO_CANCELACION_COL not in df_pedidos.columns:
        df_pedidos[MOTIVO_RECHAZO_CANCELACION_COL] = ""
    if ESTADO_ENTREGA_COL not in df_pedidos.columns:
        df_pedidos[ESTADO_ENTREGA_COL] = ""
    if df_pedidos.empty:
        st.warning("⚠️ No se pudieron cargar pedidos. Usa “🔄 Recargar…” o intenta en unos segundos.")
        if st.button("🔁 Reintentar conexión", key="retry_pedidos_inicial"):
            if allow_refresh("pedidos_last_refresh"):
                st.session_state["pedidos_reload_nonce"] += 1
                df_pedidos, headers = cargar_pedidos_desde_google_sheet(
                    GOOGLE_SHEET_ID, "datos_pedidos", st.session_state["pedidos_reload_nonce"]
                )
                st.session_state.df_pedidos = df_pedidos
                st.session_state.headers = headers
                st.toast("Reintentando...", icon="🔄")
                rerun_current_tab()
        # No st.stop(): deja que otras pestañas/partes sigan funcionando
    if FECHA_CONFIRMADO_COL not in df_pedidos.columns:
        df_pedidos[FECHA_CONFIRMADO_COL] = ""
    st.session_state.df_pedidos = df_pedidos
    st.session_state.headers = headers
    refresh_pedidos_pagados_no_confirmados(df_pedidos)

df_pedidos = st.session_state.df_pedidos
headers = st.session_state.headers
pedidos_pagados_no_confirmados = st.session_state.get('pedidos_pagados_no_confirmados', pd.DataFrame())
if not pedidos_pagados_no_confirmados.empty and FECHA_CONFIRMADO_COL not in pedidos_pagados_no_confirmados.columns:
    pedidos_pagados_no_confirmados[FECHA_CONFIRMADO_COL] = ""
if not pedidos_pagados_no_confirmados.empty:
    pedidos_pagados_no_confirmados = pedidos_pagados_no_confirmados.drop(
        columns=["display_label"], errors="ignore"
    )
    st.session_state.pedidos_pagados_no_confirmados = pedidos_pagados_no_confirmados

df_casos, headers_casos = cargar_pedidos_desde_google_sheet(GOOGLE_SHEET_ID, "casos_especiales")



# --- CONFIGURACIÓN DE AWS S3 ---
try:
    AWS_ACCESS_KEY_ID = st.secrets["aws_access_key_id"]
    AWS_SECRET_ACCESS_KEY = st.secrets["aws_secret_access_key"]
    AWS_REGION_NAME = st.secrets["aws_region"]
    S3_BUCKET_NAME = st.secrets["s3_bucket_name"]
except KeyError as e:
    st.error(f"❌ Error: Las credenciales de AWS S3 no se encontraron en Streamlit secrets. Asegúrate de que las claves 'aws_access_key_id', 'aws_secret_access_key', 'aws_region' y 's3_bucket_name' estén directamente en tus secretos de Streamlit. Clave faltante: {e}")
    st.stop() # Detiene la ejecución de la app si no se encuentran las credenciales

S3_ATTACHMENT_PREFIX = 'adjuntos_pedidos/'


def _coerce_secret_bool(value) -> bool:
    """Normaliza valores booleanos provenientes de secrets.toml."""
    if isinstance(value, bool):
        return value
    if isinstance(value, numbers.Number):
        return value != 0
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "on", "y", "si", "sí"}
    return False


S3_PUBLIC_BASE_URL = str(st.secrets.get("s3_public_base_url", "") or "").strip().rstrip("/")

if not S3_PUBLIC_BASE_URL and S3_BUCKET_NAME and AWS_REGION_NAME:
    # Fallback al dominio por defecto del bucket. Requiere que los objetos sean públicos.
    S3_PUBLIC_BASE_URL = f"https://{S3_BUCKET_NAME}.s3.{AWS_REGION_NAME}.amazonaws.com"

S3_PUBLIC_BASE_URL = S3_PUBLIC_BASE_URL.rstrip("/")

S3_USE_PERMANENT_URLS = _coerce_secret_bool(st.secrets.get("s3_use_permanent_urls", False))

if S3_PUBLIC_BASE_URL:
    S3_USE_PERMANENT_URLS = True

st.title("👨‍💼 App de Administración TD")
st.write("Panel de administración para revisar y confirmar comprobantes de pago.")

# --- FUNCIONES DE CARGA DE DATOS Y S3 (Adaptadas) ---

@st.cache_resource
def get_s3_client_cached(): # Renombrado para evitar conflicto con la variable global s3_client
    try:
        s3 = boto3.client(
            's3',
            aws_access_key_id=AWS_ACCESS_KEY_ID,
            aws_secret_access_key=AWS_SECRET_ACCESS_KEY,
            region_name=AWS_REGION_NAME
        )
        return s3
    except Exception as e:
        st.error(f"❌ Error al autenticar AWS S3: {e}")
        return None

def find_pedido_subfolder_prefix(s3_client_instance, parent_prefix, folder_name): # Acepta s3_client_instance
    if not s3_client_instance:
        return None
    
    possible_prefixes = [
        f"{parent_prefix}{folder_name}/",
        f"{parent_prefix}{folder_name}",
        f"adjuntos_pedidos/{folder_name}/",
        f"adjuntos_pedidos/{folder_name}",
        f"{folder_name}/",
        folder_name
    ]
    
    for pedido_prefix in possible_prefixes:
        try:
            response = s3_client_instance.list_objects_v2(
                Bucket=S3_BUCKET_NAME,
                Prefix=pedido_prefix,
                MaxKeys=1
            )
            
            if 'Contents' in response and response['Contents']:
                return pedido_prefix
            
        except Exception:
            continue
    
    try:
        response = s3_client_instance.list_objects_v2( # Usa s3_client_instance
            Bucket=S3_BUCKET_NAME,
            MaxKeys=100
        )
        
        if 'Contents' in response:
            for obj in response['Contents']:
                if folder_name in obj['Key']:
                    if '/' in obj['Key']:
                        prefix_parts = obj['Key'].split('/')[:-1]
                        return '/'.join(prefix_parts) + '/'
            
    except Exception:
        pass
    
    return None

def get_files_in_s3_prefix(s3_client_instance, prefix): # Acepta s3_client_instance
    if not s3_client_instance or not prefix:
        return []

    files = []
    continuation_token = None

    try:
        while True:
            request_args = {
                "Bucket": S3_BUCKET_NAME,
                "Prefix": prefix,
                "MaxKeys": 1000,
            }
            if continuation_token:
                request_args["ContinuationToken"] = continuation_token

            response = s3_client_instance.list_objects_v2(**request_args)  # Usa s3_client_instance

            for item in response.get('Contents', []):
                if not item['Key'].endswith('/'):
                    file_name = item['Key'].split('/')[-1]
                    if file_name:
                        files.append({
                            'title': file_name,
                            'key': item['Key'],
                            'size': item['Size'],
                            'last_modified': item['LastModified']
                        })

            if response.get('IsTruncated'):
                continuation_token = response.get('NextContinuationToken')
            else:
                break

        return files

    except Exception as e:
        st.error(f"❌ Error al obtener archivos del prefijo S3 '{prefix}': {e}")
        return []

def get_s3_file_download_url(s3_client_instance, object_key): # Acepta s3_client_instance
    if not object_key:
        return "#"

    key_path = str(object_key).lstrip("/")
    if not key_path:
        return "#"

    if S3_PUBLIC_BASE_URL:
        return f"{S3_PUBLIC_BASE_URL}/{key_path}"

    st.error("❌ No se pudo construir una URL pública de S3 porque falta la configuración base.")
    return "#"


def clasificar_archivos_adjuntos(files: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    """Clasifica archivos en comprobantes, facturas y otros.

    - Las imágenes se consideran comprobantes.
    - Los PDFs que no son guías se consideran facturas, aun si no dicen "factura".
    - Los archivos con "comprobante" se clasifican como comprobantes aunque no sean imagen.
    """
    comprobantes: list[dict] = []
    facturas: list[dict] = []
    otros: list[dict] = []
    image_exts = {".jpg", ".jpeg", ".png", ".gif", ".heic", ".webp"}

    for file in files:
        title = file.get("title", "")
        title_lower = title.lower()
        ext = os.path.splitext(title_lower)[1]
        is_image = ext in image_exts
        is_guia = "guia" in title_lower or "guía" in title_lower

        if "comprobante" in title_lower or is_image:
            comprobantes.append(file)
            continue

        if "factura" in title_lower or (ext == ".pdf" and not is_guia):
            facturas.append(file)
            continue

        otros.append(file)

    return comprobantes, facturas, otros
    
def _is_acl_not_supported_error(error: Exception) -> bool:
    """Return True if the boto error corresponds to AccessControlListNotSupported."""
    if not isinstance(error, ClientError):
        return False
    error_info = error.response.get("Error", {}) if hasattr(error, "response") else {}
    return error_info.get("Code") == "AccessControlListNotSupported"


def upload_file_to_s3(s3_client, bucket_name, file_obj, s3_key):
    """
    Uploads a file-like object to S3.
    Returns (success: bool, url: str)
    """
    try:
        file_obj.seek(0)  # Rebobina el archivo para iniciar la carga desde el inicio
        extra_args = {"ACL": "public-read"} if S3_USE_PERMANENT_URLS else None
        if extra_args:
            try:
                s3_client.upload_fileobj(file_obj, bucket_name, s3_key, ExtraArgs=extra_args)
            except ClientError as err:
                if _is_acl_not_supported_error(err):
                    file_obj.seek(0)
                    s3_client.upload_fileobj(file_obj, bucket_name, s3_key)
                else:
                    raise
        else:
            s3_client.upload_fileobj(file_obj, bucket_name, s3_key)
        url = get_s3_file_download_url(s3_client, s3_key)
        return True, url
    except Exception as e:
        print(f"Error uploading to S3: {e}")
        return False, ""


def build_and_upload_comprobante_index_html(
    pedido_id: str,
    comprobantes: list[dict],
    s3_client_instance,
):
    """Genera un índice HTML para múltiples comprobantes y devuelve la URL subida en S3."""
    if not pedido_id or not comprobantes or not s3_client_instance:
        return None

    try:
        items_html: list[str] = []
        for comprobante in comprobantes:
            key = comprobante.get("key")
            if not key:
                continue

            archivo_url = get_s3_file_download_url(s3_client_instance, key)

            titulo = comprobante.get("title") or key.split("/")[-1]
            titulo = html.escape(str(titulo))
            archivo_url = html.escape(str(archivo_url), quote=True)
            items_html.append(
                f'<li><a href="{archivo_url}" target="_blank" rel="noopener noreferrer">{titulo}</a></li>'
            )

        if not items_html:
            return None

        pedido_label = html.escape(str(pedido_id))
        html_content = """<!DOCTYPE html>
<html lang=\"es\">
<head>
    <meta charset=\"utf-8\" />
    <title>Comprobantes {pedido}</title>
</head>
<body>
    <h1>Comprobantes del pedido {pedido}</h1>
    <ul>
        {items}
    </ul>
</body>
</html>
""".format(pedido=pedido_label, items="\n        ".join(items_html))

        index_key = f"{S3_ATTACHMENT_PREFIX}{pedido_id}/comprobantes/index.html"
        put_kwargs = {
            "Bucket": S3_BUCKET_NAME,
            "Key": index_key,
            "Body": html_content.encode("utf-8"),
            "ContentType": "text/html",
        }
        if S3_USE_PERMANENT_URLS:
            put_kwargs["ACL"] = "public-read"

        try:
            s3_client_instance.put_object(**put_kwargs)
        except ClientError as err:
            acl_value = put_kwargs.pop("ACL", None)
            if acl_value and _is_acl_not_supported_error(err):
                s3_client_instance.put_object(**put_kwargs)
            else:
                if acl_value:
                    put_kwargs["ACL"] = acl_value
                raise

        index_url = get_s3_file_download_url(s3_client_instance, index_key)

        return index_url
    except Exception as e:
        st.warning(f"⚠️ No se pudo generar el índice de comprobantes para {pedido_id}: {e}")
        return None


def build_and_upload_comprobante_index_from_urls(
    pedido_id: str,
    urls: list[str],
    s3_client_instance,
    *,
    category: str = "comprobantes",
    page_title: str | None = None,
) -> str | None:
    """Genera y publica un índice HTML en S3 a partir de URLs externas."""

    if not s3_client_instance or not urls:
        return None

    cleaned_items: list[tuple[str, str]] = []
    for idx, raw_url in enumerate(urls, start=1):
        url_text = str(raw_url or "").strip()
        if not url_text:
            continue
        parsed = urlparse(url_text)
        filename = os.path.basename(parsed.path.rstrip("/")) if parsed.path else ""
        filename = unquote(filename) if filename else ""
        if not filename:
            filename = f"Comprobante {idx}"
        safe_title = html.escape(filename.strip() or f"Comprobante {idx}")
        safe_url = html.escape(url_text, quote=True)
        cleaned_items.append((safe_title, safe_url))

    if not cleaned_items:
        return None

    pedido_key = normalize_id_pedido(pedido_id) or str(pedido_id or "").strip()
    if not pedido_key:
        pedido_key = f"pedido-{uuid.uuid4().hex}"

    pedido_label = html.escape(str(pedido_key))
    safe_category = str(category or "comprobantes").strip().lower()
    safe_category = safe_category.replace(" ", "-").replace("/", "-")
    safe_category = safe_category or "comprobantes"
    display_title = page_title or safe_category.replace("-", " ").replace("_", " ").title()
    items_html = "\n        ".join(
        f'<li><a href="{url}" target="_blank" rel="noopener noreferrer">{title}</a></li>'
        for title, url in cleaned_items
    )

    html_content = """<!DOCTYPE html>
<html lang=\"es\">
<head>
    <meta charset=\"utf-8\" />
    <title>{display} {pedido}</title>
</head>
<body>
    <h1>{display} del pedido {pedido}</h1>
    <ul>
        {items}
    </ul>
</body>
</html>
""".format(pedido=pedido_label, items=items_html, display=html.escape(display_title))

    index_key = f"{S3_ATTACHMENT_PREFIX}{pedido_key}/{safe_category}/adjuntos-index-{uuid.uuid4().hex}.html"

    try:
        put_kwargs = {
            "Bucket": S3_BUCKET_NAME,
            "Key": index_key,
            "Body": html_content.encode("utf-8"),
            "ContentType": "text/html",
        }
        if S3_USE_PERMANENT_URLS:
            put_kwargs["ACL"] = "public-read"

        try:
            s3_client_instance.put_object(**put_kwargs)
        except ClientError as err:
            acl_value = put_kwargs.pop("ACL", None)
            if acl_value and _is_acl_not_supported_error(err):
                s3_client_instance.put_object(**put_kwargs)
            else:
                if acl_value:
                    put_kwargs["ACL"] = acl_value
                raise

        index_url = get_s3_file_download_url(s3_client_instance, index_key)

        return index_url
    except Exception as e:
        st.warning(
            f"⚠️ No se pudo generar el índice de comprobantes manual para {pedido_key}: {e}"
        )
        return None


def discover_comprobante_assets(
    pedido_id: str,
    tipo_envio: str | None,
    s3_client_instance,
) -> dict:
    """Obtiene los enlaces relevantes en S3 para un pedido y reconstruye los índices cuando aplica."""

    result: dict = {
        "files": [],
        "comprobantes": [],
        "comprobante_urls": [],
        "comprobante_link": "",
        "factura_url": "",
        "guia_url": "",
        "refact_url": "",
    }

    if not pedido_id or not s3_client_instance:
        return result

    prefix = f"{S3_ATTACHMENT_PREFIX}{pedido_id}/"
    files = get_files_in_s3_prefix(s3_client_instance, prefix)

    comprobantes = [
        f for f in files if "comprobante" in str(f.get("title", "")).lower()
    ] if files else []

    if not files or not comprobantes:
        original_prefix = find_pedido_subfolder_prefix(
            s3_client_instance, S3_ATTACHMENT_PREFIX, pedido_id
        )
        if original_prefix:
            original_files = get_files_in_s3_prefix(s3_client_instance, original_prefix)
            if original_files:
                if not files:
                    files = original_files
                else:
                    combined_files = list(files)
                    existing_keys = {
                        f.get("key")
                        for f in files
                        if isinstance(f, dict) and f.get("key")
                    }
                    for extra_file in original_files:
                        if not isinstance(extra_file, dict):
                            continue
                        extra_key = extra_file.get("key")
                        if extra_key and extra_key in existing_keys:
                            continue
                        if not extra_key and extra_file in combined_files:
                            continue
                        combined_files.append(extra_file)
                        if extra_key:
                            existing_keys.add(extra_key)
                    files = combined_files

                comprobantes = [
                    f for f in files if "comprobante" in str(f.get("title", "")).lower()
                ]

    result["files"] = files
    if not files:
        return result

    # Recalcula comprobantes a partir de la lista final para asegurar que el índice considere todos los elementos
    comprobantes = [
        f for f in files if "comprobante" in str(f.get("title", "")).lower()
    ]
    result["comprobantes"] = comprobantes

    comprobante_urls: list[str] = []
    for comprobante in comprobantes:
        key = comprobante.get("key")
        if not key:
            continue
        url = get_s3_file_download_url(s3_client_instance, key)
        if url:
            comprobante_urls.append(url)

    result["comprobante_urls"] = comprobante_urls

    if len(comprobante_urls) == 1:
        result["comprobante_link"] = comprobante_urls[0]
    elif len(comprobante_urls) > 1:
        index_url = build_and_upload_comprobante_index_html(pedido_id, comprobantes, s3_client_instance)
        if index_url and index_url != "#":
            result["comprobante_link"] = index_url
        else:
            result["comprobante_link"] = ", ".join(comprobante_urls)

    print(
        "[discover_comprobante_assets]",
        json.dumps(
            {
                "pedido_id": pedido_id,
                "total_archivos": len(files),
                "total_comprobantes": len(comprobantes),
                "comprobante_link": result.get("comprobante_link", ""),
            },
            ensure_ascii=False,
        ),
    )

    facturas = [
        f for f in files if "factura" in str(f.get("title", "")).lower()
    ]
    if facturas:
        result["factura_url"] = get_s3_file_download_url(
            s3_client_instance, facturas[0]["key"]
        )

    tipo_envio_text = str(tipo_envio or "").strip().lower()
    tipo_envio_text_flat = tipo_envio_text.replace("ó", "o")
    if "foráneo" in tipo_envio_text or "foraneo" in tipo_envio_text_flat:
        guias_filtradas = [
            f
            for f in files
            if str(f.get("title", "")).lower().endswith(".pdf")
            and re.search(r"(gu[ií]a|descarga)", str(f.get("title", "")).lower())
        ]
    else:
        guias_filtradas = [
            f for f in files if str(f.get("title", "")).lower().endswith(".xlsx")
        ]

    if guias_filtradas:
        guias_con_surtido = [
            f for f in guias_filtradas if "surtido" in str(f.get("title", "")).lower()
        ]
        guia_final = guias_con_surtido[0] if guias_con_surtido else guias_filtradas[0]
        result["guia_url"] = get_s3_file_download_url(
            s3_client_instance, guia_final["key"]
        )

    refacturas = [
        f for f in files if "surtido_factura" in str(f.get("title", "")).lower()
    ]
    if refacturas:
        result["refact_url"] = get_s3_file_download_url(
            s3_client_instance, refacturas[0]["key"]
        )

    return result

# --- Inicializar clientes de Gspread y S3 ---
try:
    gc = get_google_sheets_client()
    s3_client = get_s3_client_cached() # Ahora llama a la función cacheada
    
    if not s3_client:
        st.error("❌ No se pudo inicializar el cliente de AWS S3.")
        st.stop()
    
except Exception as e:
    st.error(f"❌ Error al autenticarse o inicializar clientes de Google Sheets/AWS S3: {e}")
    st.info("ℹ️ Asegúrate de que:")
    st.info("- Tus credenciales de Google Sheets ('google_credentials') sean correctas en secrets.toml")
    st.info("- Las APIs de Drive/Sheets estén habilitadas en Google Cloud")
    st.info("- La cuenta de servicio de Google tenga permisos en el Sheet")
    st.info("- Tus credenciales de AWS S3 (aws_access_key_id, aws_secret_access_key, aws_region) y el s3_bucket_name sean correctos en secrets.toml.")
    st.info("- La cuenta de AWS tenga permisos de lectura en el bucket S3.")
    st.stop()

# Calcular pedidos pendientes para usar en ambos tabs
pedidos_pagados_no_confirmados = st.session_state.get('pedidos_pagados_no_confirmados', pd.DataFrame())

# ---- TABS ADMIN ----
# Mantiene la pestaña activa usando los query params de Streamlit
tab_names = [
    "💳 Pendientes de Confirmar",
    "📥 Confirmados",
    "📦 Casos Especiales",
    "🗂️ Data Especiales",
]

_tab_param = st.query_params.get(TAB_QUERY_PARAM)
if isinstance(_tab_param, (list, tuple)):
    _tab_param = _tab_param[0] if _tab_param else None

if _tab_param is None:
    _session_tab_raw = st.session_state.get(TAB_SESSION_KEY)
    if _session_tab_raw is None:
        _session_tab_raw = st.session_state.get("current_tab")
else:
    _session_tab_raw = _tab_param

_default_tab = _normalize_tab_index(_session_tab_raw, len(tab_names))

st.session_state[TAB_SESSION_KEY] = _default_tab
st.session_state["current_tab"] = str(_default_tab)

if st.query_params.get(TAB_QUERY_PARAM) != str(_default_tab):
    st.query_params[TAB_QUERY_PARAM] = str(_default_tab)

tabs = st.tabs(tab_names)
tab1, tab2, tab3, tab4 = tabs

tab1_is_active = _default_tab == 0
tab2_is_active = _default_tab == 1
tab3_is_active = _default_tab == 2
tab4_is_active = _default_tab == 3

# forza la pestaña almacenada al recargar y sincroniza la URL al cambiar de pestaña
components.html(
    f"""
    <script>
    (function() {{
        const desiredIndex = {_default_tab};
        const parentWindow = window.parent;
        const parentDocument = parentWindow.document;

        function updateQueryParam(index) {{
            try {{
                const url = new URL(parentWindow.location.href);
                url.searchParams.set('{TAB_QUERY_PARAM}', index);
                parentWindow.history.replaceState(null, '', url.toString());
            }} catch (error) {{
                console.error('Error syncing tab param', error);
            }}
        }}

        function attachListeners() {{
            const buttons = parentDocument.querySelectorAll('div[data-baseweb="tab-list"] button');
            if (!buttons || !buttons.length) {{
                setTimeout(attachListeners, 100);
                return;
            }}

            buttons.forEach((button, index) => {{
                if (button.getAttribute('data-tab-listener') === 'true') {{
                    return;
                }}
                button.setAttribute('data-tab-listener', 'true');
                button.addEventListener('click', () => updateQueryParam(String(index)));
            }});

            if (desiredIndex >= 0 && desiredIndex < buttons.length) {{
                const targetButton = buttons[desiredIndex];
                if (targetButton && targetButton.getAttribute('aria-selected') !== 'true') {{
                    targetButton.click();
                }}
            }}
        }}

        if (document.readyState === 'complete') {{
            attachListeners();
        }} else {{
            window.addEventListener('load', attachListeners);
            document.addEventListener('DOMContentLoaded', attachListeners);
            setTimeout(attachListeners, 500);
        }}
    }})();
    </script>
    """,
    height=0,
)


# --- INTERFAZ PRINCIPAL ---
with tab1:
    if tab1_is_active:
        st.session_state[TAB_SESSION_KEY] = 0
        st.session_state["current_tab"] = "0"
    st.header("💳 Comprobantes de Pago Pendientes de Confirmación")
    mostrar = True  # ✅ Se inicializa desde el inicio del tab

    # Limpia la bandera de auto-recarga heredada de versiones anteriores
    st.session_state.pop("pedidos_autorefresh", None)

    if st.button("🔄 Recargar Pedidos desde Google Sheets", type="secondary"):
        if allow_refresh("pedidos_last_refresh"):
            # Borra el snapshot cacheado para forzar que la lectura venga fresca de Google Sheets.
            cargar_pedidos_desde_google_sheet.clear()
            st.session_state["pedidos_reload_nonce"] += 1
            df_pedidos, headers = cargar_pedidos_desde_google_sheet(
                GOOGLE_SHEET_ID, "datos_pedidos", st.session_state["pedidos_reload_nonce"]
            )
            if 'Tipo_Envio' in df_pedidos.columns:
                df_pedidos = df_pedidos[
                    ~df_pedidos['Tipo_Envio'].isin(['🎓 Cursos y Eventos', '📋 Solicitudes de Guía'])
                ].copy()
            if MOTIVO_RECHAZO_CANCELACION_COL not in df_pedidos.columns:
                df_pedidos[MOTIVO_RECHAZO_CANCELACION_COL] = ""
            if ESTADO_ENTREGA_COL not in df_pedidos.columns:
                df_pedidos[ESTADO_ENTREGA_COL] = ""
            st.session_state.df_pedidos = df_pedidos
            st.session_state.headers = headers
            refresh_pedidos_pagados_no_confirmados(df_pedidos)
            st.toast("Pedidos recargados", icon="🔄")

    if df_pedidos.empty:
        st.info("ℹ️ No hay pedidos cargados en este momento.")
    else:
        if pedidos_pagados_no_confirmados.empty:
            st.success("🎉 ¡No hay comprobantes pendientes de confirmación!")
            st.info("Todos los pedidos pagados han sido confirmados.")
        else:
            st.warning(f"📋 Hay {len(pedidos_pagados_no_confirmados)} comprobantes pendientes.")

            # Mostrar tabla
            columns_to_show = [
                'Folio_Factura', 'Cliente', 'Vendedor_Registro', 'Tipo_Envio',
                'Fecha_Entrega', 'Estado', 'Estado_Pago'
            ]
            turno_display_col = "Turno (Local)"
            has_turno_data = "Turno" in pedidos_pagados_no_confirmados.columns

            if has_turno_data:
                try:
                    insert_idx = columns_to_show.index("Tipo_Envio") + 1
                except ValueError:
                    insert_idx = len(columns_to_show)
                columns_to_show.insert(insert_idx, turno_display_col)

            existing_columns: list[str] = []
            for col in columns_to_show:
                if col == turno_display_col:
                    if has_turno_data:
                        existing_columns.append(col)
                elif col in pedidos_pagados_no_confirmados.columns:
                    existing_columns.append(col)

            if existing_columns:
                base_columns = [c for c in existing_columns if c != turno_display_col]
                df_vista = pedidos_pagados_no_confirmados[base_columns].copy()

                if has_turno_data:
                    def _format_turno_display(row: pd.Series) -> str:
                        tipo_envio = str(row.get("Tipo_Envio", "")).strip()
                        if tipo_envio != "📍 Pedido Local":
                            return ""

                        turno_raw = str(row.get("Turno", "") or "").strip()
                        if not turno_raw or turno_raw.lower() in {"nan", "none"}:
                            return ""

                        return turno_raw

                    df_vista[turno_display_col] = pedidos_pagados_no_confirmados.apply(
                        _format_turno_display, axis=1
                    )
                    ordered_cols = [
                        col for col in existing_columns if col in df_vista.columns
                    ]
                    df_vista = df_vista[ordered_cols]

                if 'Fecha_Entrega' in df_vista.columns:
                    df_vista['Fecha_Entrega'] = pd.to_datetime(df_vista['Fecha_Entrega'], errors='coerce').dt.strftime('%d/%m/%Y')

                st.dataframe(
                    df_vista.sort_values(by='Fecha_Entrega' if 'Fecha_Entrega' in df_vista.columns else existing_columns[0]),
                    use_container_width=True,
                    hide_index=True
                )

            st.markdown("---")
            st.subheader("🔍 Revisar Comprobante de Pago")

            # Opciones de selección
            pedidos_pagados_no_confirmados['display_label'] = pedidos_pagados_no_confirmados.apply(lambda row: (
                f"📄 {row.get('Folio_Factura', 'N/A')} - "
                f"👤 {row.get('Cliente', 'N/A')} - "
                f"{row.get('Estado', 'N/A')} - "
                f"{row.get('Tipo_Envio', 'N/A')}"
            ), axis=1)

            pedido_options = pedidos_pagados_no_confirmados['display_label'].tolist()
            selected_index = st.selectbox(
                "📝 Seleccionar Pedido para Revisar Comprobante",
                options=range(len(pedido_options)),
                format_func=lambda i: pedido_options[i],
                key="select_pedido_comprobante"
            )

            if selected_index is not None:
                selected_pedido_data = pedidos_pagados_no_confirmados.iloc[selected_index]

                raw_pedido_id = selected_pedido_data.get("ID_Pedido")
                normalized_pedido_id = normalize_id_pedido(raw_pedido_id)
                fallback_folio = normalize_folio_factura(
                    selected_pedido_data.get("Folio_Factura")
                )

                if normalized_pedido_id:
                    current_selection_key = f"ID::{normalized_pedido_id}"
                elif fallback_folio:
                    current_selection_key = f"FOLIO::{fallback_folio}"
                else:
                    current_selection_key = f"INDEX::{selected_index}"

                last_selection_key = st.session_state.get("last_selected_pedido_key")
                if last_selection_key != current_selection_key:
                    clear_comprobante_form_state()

                st.session_state["last_selected_pedido_key"] = current_selection_key
                modificacion_surtido_text = clean_modificacion_surtido(
                    selected_pedido_data.get("Modificacion_Surtido", "")
                )

                is_pedido_local = (
                    str(selected_pedido_data.get("Tipo_Envio", "")).strip()
                    == "📍 Pedido Local"
                )
                estado_entrega_stored = normalize_estado_entrega(
                    selected_pedido_data.get(ESTADO_ENTREGA_COL, "")
                )
                if estado_entrega_stored not in ESTADO_ENTREGA_OPCIONES:
                    estado_entrega_stored = ESTADO_ENTREGA_DEFAULT
                estado_entrega_index = ESTADO_ENTREGA_OPCIONES.index(estado_entrega_stored)
                estado_entrega_value = estado_entrega_stored
                estado_entrega_widget_key = _comprobante_form_key("estado_entrega_local")
                id_vendedor_user = normalize_user_field(
                    selected_pedido_data.get("id_vendedor")
                )
                id_vendedor_mod_user = normalize_user_field(
                    selected_pedido_data.get("id_vendedor_Mod")
                )

                # 🚨 Lógica especial si es pedido a crédito
                if selected_pedido_data.get("Estado_Pago", "").strip() == "💳 CREDITO":
                    st.subheader("📝 Confirmación de Pedido a Crédito")
                    selected_pedido_id_for_s3_search = selected_pedido_data.get('ID_Pedido', 'N/A')
                    st.session_state.selected_admin_pedido_id = selected_pedido_id_for_s3_search

                    # Mostrar información del pedido
                    col1, col2 = st.columns(2)
                    with col1:
                        st.subheader("📋 Información del Pedido")
                        if modificacion_surtido_text:
                            st.info(f"🛠 Modificación de surtido: {modificacion_surtido_text}")
                        st.write(f"**📄 Folio Factura:** {selected_pedido_data.get('Folio_Factura', 'N/A')}")
                        st.write(f"**🗒 Comentario del Pedido:** {selected_pedido_data.get('Comentario', 'Sin comentario')}")
                        st.write(f"**🤝 Cliente:** {selected_pedido_data.get('Cliente', 'N/A')}")
                        st.write(f"**🧑‍💼 Vendedor:** {selected_pedido_data.get('Vendedor_Registro', 'N/A')}")
                        if id_vendedor_user:
                            st.write(f"**Usuario:** {id_vendedor_user}")
                        if id_vendedor_mod_user:
                            st.write(f"**Usuario Mod:** {id_vendedor_mod_user}")
                        st.write(f"**Tipo de Envío:** {selected_pedido_data.get('Tipo_Envio', 'N/A')}")
                        st.write(f"**📅 Fecha de Entrega:** {selected_pedido_data.get('Fecha_Entrega', 'N/A')}")
                        st.write(f"**Estado:** {selected_pedido_data.get('Estado', 'N/A')}")
                        st.write(f"**Estado de Pago:** {selected_pedido_data.get('Estado_Pago', 'N/A')}")
                        if is_pedido_local:
                            estado_entrega_value = st.selectbox(
                                "🚚 Estado de entrega",
                                ESTADO_ENTREGA_OPCIONES,
                                index=estado_entrega_index,
                                key=estado_entrega_widget_key,
                            )

                    with col2:
                        st.subheader("📎 Archivos y Comprobantes")
                        if s3_client:
                            pedido_folder_prefix = find_pedido_subfolder_prefix(s3_client, S3_ATTACHMENT_PREFIX, selected_pedido_id_for_s3_search)
                            files = get_files_in_s3_prefix(s3_client, pedido_folder_prefix) if pedido_folder_prefix else []

                            if files:
                                comprobantes, facturas, otros = clasificar_archivos_adjuntos(files)

                                if comprobantes:
                                    st.write("**🧾 Comprobantes de Pago:**")
                                    for f in comprobantes:
                                        url = get_s3_file_download_url(s3_client, f['key'])
                                        nombre = f['title'].replace(selected_pedido_id_for_s3_search, "").strip("_-")
                                        st.markdown(f"- 📄 **{nombre}** ({f['size']} bytes) [🔗 Ver/Descargar]({url})")
                                else:
                                    st.warning("⚠️ No se encontraron comprobantes.")

                                if facturas:
                                    st.write("**📑 Facturas de Venta:**")
                                    for f in facturas:
                                        url = get_s3_file_download_url(s3_client, f['key'])
                                        nombre = f['title'].replace(selected_pedido_id_for_s3_search, "").strip("_-")
                                        st.markdown(f"- 📄 **{nombre}** ({f['size']} bytes) [🔗 Ver/Descargar]({url})")

                                if otros:
                                    with st.expander("📂 Otros archivos del pedido"):
                                        for f in otros:
                                            url = get_s3_file_download_url(s3_client, f['key'])
                                            st.markdown(f"- 📄 **{f['title']}** ({f['size']} bytes) [🔗 Ver/Descargar]({url})")
                            else:
                                st.info("📁 No se encontraron archivos en la carpeta del pedido.")
                        else:
                            st.error("❌ Error de conexión con S3. Revisa las credenciales.")

                    # Confirmación de crédito
                    confirmacion_credito = st.selectbox("¿Confirmar que el pedido fue autorizado como crédito?", ["", "Sí", "No"])
                    comentario_credito = st.text_area("✍️ Comentario administrativo")

                    if confirmacion_credito:
                        if st.button("💾 Guardar Confirmación de Crédito"):
                            try:
                                # Índice real (fila en Google Sheets)
                                gsheet_row_index = df_pedidos[df_pedidos['ID_Pedido'] == selected_pedido_data["ID_Pedido"]].index[0] + 2

                                # 🔹 OBTENER HOJA FRESCA (con reintentos) ANTES DE ESCRIBIR
                                worksheet = _get_ws_datos()
                                headers = ensure_sheet_column(worksheet, headers, FECHA_CONFIRMADO_COL)
                                if is_pedido_local:
                                    headers = ensure_sheet_column(worksheet, headers, ESTADO_ENTREGA_COL)
                                st.session_state.headers = headers

                                # Actualizaciones
                                updates = []
                                local_updates = {}
                                if FECHA_CONFIRMADO_COL not in df_pedidos.columns:
                                    df_pedidos[FECHA_CONFIRMADO_COL] = ""
                                fecha_confirmado = (
                                    obtener_fecha_confirmado_cdmx()
                                    if confirmacion_credito == "Sí"
                                    else ""
                                )
                                if "Comprobante_Confirmado" in headers:
                                    updates.append({
                                        "range": rowcol_to_a1(
                                            gsheet_row_index,
                                            headers.index("Comprobante_Confirmado") + 1,
                                        ),
                                        "values": [[confirmacion_credito]],
                                    })
                                    local_updates["Comprobante_Confirmado"] = confirmacion_credito

                                if FECHA_CONFIRMADO_COL in headers:
                                    updates.append({
                                        "range": rowcol_to_a1(
                                            gsheet_row_index,
                                            headers.index(FECHA_CONFIRMADO_COL) + 1,
                                        ),
                                        "values": [[fecha_confirmado]],
                                    })
                                    local_updates[FECHA_CONFIRMADO_COL] = fecha_confirmado

                                if is_pedido_local and ESTADO_ENTREGA_COL in headers:
                                    updates.append({
                                        "range": rowcol_to_a1(
                                            gsheet_row_index,
                                            headers.index(ESTADO_ENTREGA_COL) + 1,
                                        ),
                                        "values": [[estado_entrega_value]],
                                    })
                                    local_updates[ESTADO_ENTREGA_COL] = estado_entrega_value

                                if "Comentario" in headers:
                                    comentario_existente = selected_pedido_data.get("Comentario", "")
                                    nuevo_comentario = f"Comentario de CREDITO: {comentario_credito.strip()}"
                                    comentario_final = (
                                        f"{comentario_existente}\n{nuevo_comentario}"
                                        if comentario_existente
                                        else nuevo_comentario
                                    )
                                    updates.append({
                                        "range": rowcol_to_a1(
                                            gsheet_row_index,
                                            headers.index("Comentario") + 1,
                                        ),
                                        "values": [[comentario_final]],
                                    })
                                    local_updates["Comentario"] = comentario_final

                                if updates:
                                    safe_batch_update(worksheet, updates)
                                    _get_ws_datos.clear()
                                    cargar_pedidos_desde_google_sheet.clear()

                                df_idx = df_pedidos[df_pedidos['ID_Pedido'] == selected_pedido_data["ID_Pedido"]].index
                                if len(df_idx) > 0:
                                    df_idx = df_idx[0]
                                    for col, val in local_updates.items():
                                        if col in df_pedidos.columns:
                                            df_pedidos.at[df_idx, col] = val
                                st.session_state.df_pedidos = df_pedidos
                                pedidos_pagados_no_confirmados = refresh_pedidos_pagados_no_confirmados(
                                    df_pedidos
                                )

                                st.success("✅ Confirmación de crédito guardada exitosamente.")
                                st.balloons()
                                rerun_current_tab()

                            except Exception as e:
                                st.error(f"❌ Error al guardar la confirmación: {e}")
                    else:
                        st.info("Selecciona una opción para confirmar el crédito.")


                    # 🚫 Importante: esta rama solo muestra la confirmación de crédito.
                    # La lógica de comprobantes estándar vive en el bloque 'else'.

                else:
                    # ✅ Continuar con lógica normal para pedidos no-crédito
                    if (
                        selected_pedido_data.get("Estado_Pago", "").strip() == "🔴 No Pagado" and
                        selected_pedido_data.get("Tipo_Envio", "").strip() == "📍 Pedido Local"
                    ):
                        st.subheader("🧾 Subir Comprobante de Pago")

                    if is_pedido_local:
                        estado_entrega_value = st.selectbox(
                            "🚚 Estado de entrega",
                            ESTADO_ENTREGA_OPCIONES,
                            index=estado_entrega_index,
                            key=estado_entrega_widget_key,
                        )

                    pago_doble = st.checkbox(
                        "✅ Pago en dos partes distintas",
                        key=_comprobante_form_key("pago_doble_admin"),
                    )
    
                    comprobantes_nuevo = []
                    if not pago_doble:
                        comprobantes_nuevo = st.file_uploader(
                            "📤 Subir Comprobante(s) de Pago",
                            type=["pdf", "jpg", "jpeg", "png"],
                            accept_multiple_files=True,
                            key=_comprobante_form_key("comprobante_local_no_pagado"),
                        )
    
                        with st.expander(
                            "📝 Detalles del Pago",
                            expanded=selected_pedido_data.get("Estado_Pago", "").strip() == "🔴 No Pagado",
                        ):
                            fecha_pago = st.date_input(
                                "📅 Fecha del Pago",
                                value=datetime.today().date(),
                                key=_comprobante_form_key("fecha_pago_local"),
                            )
                            forma_pago = st.selectbox(
                                "💳 Forma de Pago",
                                [
                                    "Transferencia",
                                    "Depósito en Efectivo",
                                    "Tarjeta de Débito",
                                    "Tarjeta de Crédito",
                                    "Cheque",
                                ],
                                key=_comprobante_form_key("forma_pago_local"),
                            )
                            monto_pago = st.number_input(
                                "💲 Monto del Pago",
                                min_value=0.0,
                                format="%.2f",
                                key=_comprobante_form_key("monto_pago_local"),
                            )
    
                            if forma_pago in ["Tarjeta de Débito", "Tarjeta de Crédito"]:
                                terminal = st.selectbox(
                                    "🏧 Terminal",
                                    [
                                        "BANORTE",
                                        "AFIRME",
                                        "VELPAY",
                                        "CLIP",
                                        "PAYPAL",
                                        "BBVA",
                                        "CONEKTA",
                                        "MERCADO PAGO",
                                    ],
                                    key=_comprobante_form_key("terminal_local"),
                                )
                                banco_destino = ""
                            else:
                                banco_destino = st.selectbox(
                                    "🏦 Banco Destino",
                                    [
                                        "BANORTE",
                                        "BANAMEX",
                                        "AFIRME",
                                        "BANCOMER OP",
                                        "BANCOMER CURSOS",
                                    ],
                                    key=_comprobante_form_key("banco_destino_local"),
                                )
                                terminal = ""

                            referencia = st.text_input(
                                "🔢 Referencia (opcional)",
                                key=_comprobante_form_key("referencia_local"),
                            )
    
                    else:
                        st.markdown("### 1️⃣ Primer Pago")
                        comp1 = st.file_uploader(
                            "💳 Comprobante 1",
                            type=["pdf", "jpg", "jpeg", "png"],
                            accept_multiple_files=True,
                            key=_comprobante_form_key("cp_pago1_admin"),
                        )
                        fecha1 = st.date_input(
                            "📅 Fecha 1",
                            value=datetime.today().date(),
                            key=_comprobante_form_key("fecha_pago1_admin"),
                        )
                        forma1 = st.selectbox(
                            "💳 Forma 1",
                            [
                                "Transferencia",
                                "Depósito en Efectivo",
                                "Tarjeta de Débito",
                                "Tarjeta de Crédito",
                                "Cheque",
                            ],
                            key=_comprobante_form_key("forma_pago1_admin"),
                        )
                        monto1 = st.number_input(
                            "💲 Monto 1",
                            min_value=0.0,
                            format="%.2f",
                            key=_comprobante_form_key("monto_pago1_admin"),
                        )
                        terminal1 = banco1 = ""
                        if forma1 in ["Tarjeta de Débito", "Tarjeta de Crédito"]:
                            terminal1 = st.selectbox(
                                "🏧 Terminal 1",
                                [
                                    "BANORTE",
                                    "AFIRME",
                                    "VELPAY",
                                    "CLIP",
                                    "PAYPAL",
                                    "BBVA",
                                    "CONEKTA",
                                    "MERCADO PAGO",
                                ],
                                key=_comprobante_form_key("terminal1_admin"),
                            )
                        else:
                            banco1 = st.selectbox(
                                "🏦 Banco 1",
                                [
                                    "BANORTE",
                                    "BANAMEX",
                                    "AFIRME",
                                    "BANCOMER OP",
                                    "BANCOMER CURSOS",
                                ],
                                key=_comprobante_form_key("banco1_admin"),
                            )
                        ref1 = st.text_input(
                            "🔢 Referencia 1",
                            key=_comprobante_form_key("ref1_admin"),
                        )
    
                        st.markdown("### 2️⃣ Segundo Pago")
                        comp2 = st.file_uploader(
                            "💳 Comprobante 2",
                            type=["pdf", "jpg", "jpeg", "png"],
                            accept_multiple_files=True,
                            key=_comprobante_form_key("cp_pago2_admin"),
                        )
                        fecha2 = st.date_input(
                            "📅 Fecha 2",
                            value=datetime.today().date(),
                            key=_comprobante_form_key("fecha_pago2_admin"),
                        )
                        forma2 = st.selectbox(
                            "💳 Forma 2",
                            [
                                "Transferencia",
                                "Depósito en Efectivo",
                                "Tarjeta de Débito",
                                "Tarjeta de Crédito",
                                "Cheque",
                            ],
                            key=_comprobante_form_key("forma_pago2_admin"),
                        )
                        monto2 = st.number_input(
                            "💲 Monto 2",
                            min_value=0.0,
                            format="%.2f",
                            key=_comprobante_form_key("monto_pago2_admin"),
                        )
                        terminal2 = banco2 = ""
                        if forma2 in ["Tarjeta de Débito", "Tarjeta de Crédito"]:
                            terminal2 = st.selectbox(
                                "🏧 Terminal 2",
                                [
                                    "BANORTE",
                                    "AFIRME",
                                    "VELPAY",
                                    "CLIP",
                                    "PAYPAL",
                                    "BBVA",
                                    "CONEKTA",
                                    "MERCADO PAGO",
                                ],
                                key=_comprobante_form_key("terminal2_admin"),
                            )
                        else:
                            banco2 = st.selectbox(
                                "🏦 Banco 2",
                                [
                                    "BANORTE",
                                    "BANAMEX",
                                    "AFIRME",
                                    "BANCOMER OP",
                                    "BANCOMER CURSOS",
                                ],
                                key=_comprobante_form_key("banco2_admin"),
                            )
                        ref2 = st.text_input(
                            "🔢 Referencia 2",
                            key=_comprobante_form_key("ref2_admin"),
                        )
    
                        # Unificar comprobantes y campos
                        comprobantes_nuevo = (comp1 or []) + (comp2 or [])
                        fecha_pago = f"{fecha1.strftime('%Y-%m-%d')} y {fecha2.strftime('%Y-%m-%d')}"
                        forma_pago = f"{forma1}, {forma2}"
                        terminal = f"{terminal1}, {terminal2}" if forma1.startswith("Tarjeta") or forma2.startswith("Tarjeta") else ""
                        banco_destino = f"{banco1}, {banco2}" if forma1 not in ["Tarjeta de Débito", "Tarjeta de Crédito"] or forma2 not in ["Tarjeta de Débito", "Tarjeta de Crédito"] else ""
                        monto_pago = monto1 + monto2
                        referencia = f"{ref1}, {ref2}"
    
                    # ⬇️ Reemplaza desde aquí
                    if st.button("💾 Guardar Comprobante y Datos de Pago"):
                        try:
                            # Índice real en la hoja
                            gsheet_row_index = (
                                df_pedidos[df_pedidos['ID_Pedido'] == selected_pedido_data["ID_Pedido"]]
                                .index[0] + 2
                            )
    
                            # Subir archivos a S3
                            adjuntos_urls = []
                            if comprobantes_nuevo:
                                for file in comprobantes_nuevo:
                                    ext = os.path.splitext(file.name)[1]
                                    s3_key = (
                                        f"{selected_pedido_data['ID_Pedido']}/"
                                        f"comprobante_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:4]}{ext}"
                                    )
                                    try:
                                        file.seek(0)
                                    except Exception:
                                        pass
                                    success, url = upload_file_to_s3(s3_client, S3_BUCKET_NAME, file, s3_key)
                                    if success:
                                        adjuntos_urls.append(url)
                                        st.toast(f"✅ {file.name} subido correctamente", icon="✅")
    
                                if len(adjuntos_urls) == len(comprobantes_nuevo):
                                    st.success("Todos los comprobantes se han subido correctamente")
    
                            # ---- Normalizaciones SEGURAS para Google Sheets ----
                            if isinstance(fecha_pago, (datetime, date)):
                                fecha_pago_str = fecha_pago.strftime("%Y-%m-%d")
                            else:
                                fecha_pago_str = str(fecha_pago) if fecha_pago else ""
    
                            try:
                                monto_val = float(monto_pago) if monto_pago is not None else 0.0
                            except Exception:
                                monto_val = 0.0
    
                            if FECHA_CONFIRMADO_COL not in df_pedidos.columns:
                                df_pedidos[FECHA_CONFIRMADO_COL] = ""
                            fecha_confirmado = obtener_fecha_confirmado_cdmx()

                            updates = {
                                "Estado_Pago": "✅ Pagado",
                                "Comprobante_Confirmado": "Sí",
                                FECHA_CONFIRMADO_COL: fecha_confirmado,
                                "Fecha_Pago_Comprobante": fecha_pago_str,
                                "Forma_Pago_Comprobante": forma_pago,
                                "Monto_Comprobante": monto_val,
                                "Referencia_Comprobante": referencia,
                                "Terminal": terminal,
                                "Banco_Destino_Pago": banco_destino,
                            }
                            if is_pedido_local:
                                updates[ESTADO_ENTREGA_COL] = estado_entrega_value
    
                            # 🔹 OBTENER HOJA FRESCA (con reintentos) ANTES DE ESCRIBIR
                            worksheet = _get_ws_datos()
                            headers = ensure_sheet_column(worksheet, headers, FECHA_CONFIRMADO_COL)
                            if is_pedido_local:
                                headers = ensure_sheet_column(worksheet, headers, ESTADO_ENTREGA_COL)
                            st.session_state.headers = headers

                            cell_updates = []
                            nuevo_valor_adjuntos = None

                            # Escribir columnas principales
                            for col, val in updates.items():
                                if col in headers:
                                    cell_updates.append({
                                        "range": rowcol_to_a1(
                                            gsheet_row_index, headers.index(col) + 1
                                        ),
                                        "values": [[val]],
                                    })
    
                            # Concatenar nuevos adjuntos al campo "Adjuntos"
                            if adjuntos_urls and "Adjuntos" in headers:
                                adjuntos_actuales = selected_pedido_data.get("Adjuntos", "")
                                nuevo_valor_adjuntos = ", ".join(
                                    filter(None, [adjuntos_actuales] + adjuntos_urls)
                                )
                                cell_updates.append({
                                    "range": rowcol_to_a1(
                                        gsheet_row_index, headers.index("Adjuntos") + 1
                                    ),
                                    "values": [[nuevo_valor_adjuntos]],
                                })
    
    
                            if cell_updates:
                                safe_batch_update(worksheet, cell_updates)
                                _get_ws_datos.clear()
    
                            df_idx = df_pedidos[df_pedidos['ID_Pedido'] == selected_pedido_data["ID_Pedido"]].index
                            if len(df_idx) > 0:
                                df_idx = df_idx[0]
                                for col, val in updates.items():
                                    if col in df_pedidos.columns:
                                        df_pedidos.at[df_idx, col] = val
                                if nuevo_valor_adjuntos and 'Adjuntos' in df_pedidos.columns:
                                    df_pedidos.at[df_idx, 'Adjuntos'] = nuevo_valor_adjuntos
                            st.session_state.df_pedidos = df_pedidos
                            pedidos_pagados_no_confirmados = refresh_pedidos_pagados_no_confirmados(
                                df_pedidos
                            )

                            clear_comprobante_form_state()
                            st.session_state.pop("last_selected_pedido_key", None)
                            st.success("✅ Comprobante y datos de pago guardados exitosamente.")
                            st.balloons()
                            rerun_current_tab()
    
                        except Exception as e:
                            st.error(f"❌ Error al guardar el comprobante: {e}")
    
                    # ⬆️ Hasta aquí
    
    
                    # Resto del código para pedidos normales con comprobantes existentes
                    selected_pedido_id_for_s3_search = selected_pedido_data.get('ID_Pedido', 'N/A')
    
                    st.session_state.selected_admin_pedido_id = selected_pedido_id_for_s3_search
                    fecha_pago_raw = selected_pedido_data.get('Fecha_Pago_Comprobante')
                    if isinstance(fecha_pago_raw, str) and " y " in fecha_pago_raw:
                        st.session_state.fecha_pago = fecha_pago_raw  # ya viene listo como string concatenado
                    else:
                        st.session_state.fecha_pago = pd.to_datetime(fecha_pago_raw).date() if fecha_pago_raw else None
    
                    st.session_state.forma_pago = selected_pedido_data.get('Forma_Pago_Comprobante', 'Transferencia')
                    valid_terminals = ["BANORTE", "AFIRME", "VELPAY", "CLIP", "PAYPAL", "BBVA", "CONEKTA", "MERCADO PAGO"]
                    st.session_state.terminal = selected_pedido_data.get('Terminal', 'BANORTE')
                    if st.session_state.terminal not in valid_terminals:
                        st.session_state.terminal = 'BANORTE'
                    st.session_state.banco_destino_pago = selected_pedido_data.get('Banco_Destino_Pago', 'BANORTE')
                    try:
                        st.session_state.monto_pago = float(selected_pedido_data.get('Monto_Comprobante', 0.0))
                    except Exception:
                        st.session_state.monto_pago = 0.0
                    st.session_state.referencia_pago = selected_pedido_data.get('Referencia_Comprobante', '')
    
                    col1, col2 = st.columns(2)
                    with col1:
                        st.subheader("📋 Información del Pedido")
                        if modificacion_surtido_text:
                            st.info(f"🛠 Modificación de surtido: {modificacion_surtido_text}")
                        st.write(f"**Folio Factura:** {selected_pedido_data.get('Folio_Factura', 'N/A')}")
                        st.write(f"**🗒 Comentario del Pedido:** {selected_pedido_data.get('Comentario', 'Sin comentario')}")
                        st.write(f"**Cliente:** {selected_pedido_data.get('Cliente', 'N/A')}")
                        st.write(f"**Vendedor:** {selected_pedido_data.get('Vendedor_Registro', 'N/A')}")
                        if id_vendedor_user:
                            st.write(f"**Usuario:** {id_vendedor_user}")
                        if id_vendedor_mod_user:
                            st.write(f"**Usuario Mod:** {id_vendedor_mod_user}")
                        st.write(f"**Tipo de Envío:** {selected_pedido_data.get('Tipo_Envio', 'N/A')}")
                        st.write(f"**Fecha de Entrega:** {selected_pedido_data.get('Fecha_Entrega', 'N/A')}")
                        st.write(f"**Estado:** {selected_pedido_data.get('Estado', 'N/A')}")
                        st.write(f"**Estado de Pago:** {selected_pedido_data.get('Estado_Pago', 'N/A')}")
    
                    with col2:
                        st.subheader("📎 Archivos y Comprobantes")
    
                        if s3_client:
                            pedido_folder_prefix = find_pedido_subfolder_prefix(s3_client, S3_ATTACHMENT_PREFIX, selected_pedido_id_for_s3_search)
                            files = get_files_in_s3_prefix(s3_client, pedido_folder_prefix) if pedido_folder_prefix else []

                            if files:
                                comprobantes, facturas, otros = clasificar_archivos_adjuntos(files)

                                if comprobantes:
                                    st.write("**🧾 Comprobantes de Pago:**")
                                    for f in comprobantes:
                                        url = get_s3_file_download_url(s3_client, f['key'])
                                        nombre = f['title'].replace(selected_pedido_id_for_s3_search, "").strip("_-")
                                        st.markdown(f"- 📄 **{nombre}** ({f['size']} bytes) [🔗 Ver/Descargar]({url})")
                                else:
                                    st.warning("⚠️ No se encontraron comprobantes.")
    
                                if facturas:
                                    st.write("**📑 Facturas de Venta:**")
                                    for f in facturas:
                                        url = get_s3_file_download_url(s3_client, f['key'])
                                        nombre = f['title'].replace(selected_pedido_id_for_s3_search, "").strip("_-")
                                        st.markdown(f"- 📄 **{nombre}** ({f['size']} bytes) [🔗 Ver/Descargar]({url})")
    
                                if otros:
                                    with st.expander("📂 Otros archivos del pedido"):
                                        for f in otros:
                                            url = get_s3_file_download_url(s3_client, f['key'])
                                            st.markdown(f"- 📄 **{f['title']}** ({f['size']} bytes) [🔗 Ver/Descargar]({url})")
                            else:
                                st.info("📁 No se encontraron archivos en la carpeta del pedido.")
                        else:
                            st.error("❌ Error de conexión con S3. Revisa las credenciales.")
    
                    # Detectar cuántos comprobantes hay
                    num_comprobantes = len(comprobantes) if 'comprobantes' in locals() else 0
                    estado_pago_raw = str(selected_pedido_data.get("Estado_Pago", "")).strip()
                    pedido_pagado_sin_confirmar = estado_pago_raw.startswith("✅")

                    # Para pedidos marcados como pagados, permite capturar datos aunque no haya archivos en S3
                    if pedido_pagado_sin_confirmar and num_comprobantes == 0:
                        num_comprobantes = 1
                        st.info(
                            "ℹ️ El pedido ya está marcado como pagado. Puedes llenar el formulario y confirmar el comprobante aunque no existan archivos adjuntos."
                        )

                    st.subheader("✅ Confirmar Comprobante")

                    fecha_list: list[str] = []
                    forma_list: list[str] = []
                    banco_list: list[str] = []
                    terminal_list: list[str] = []
                    monto_list: list[float] = []
                    ref_list: list[str] = []

                    if num_comprobantes == 0:
                        st.warning("⚠️ No hay comprobantes para confirmar.")
                    else:
                        # --- Prellenar valores si ya están registrados en la hoja ---
                        fecha_list = str(selected_pedido_data.get('Fecha_Pago_Comprobante', '')).split(" y ")
                        forma_list = str(selected_pedido_data.get('Forma_Pago_Comprobante', '')).split(", ")
                        banco_list = str(selected_pedido_data.get('Banco_Destino_Pago', '')).split(", ")
                        terminal_list = str(selected_pedido_data.get('Terminal', '')).split(", ")
                        monto_list_raw = selected_pedido_data.get('Monto_Comprobante', '')
                        ref_list = str(selected_pedido_data.get('Referencia_Comprobante', '')).split(", ")

                        if isinstance(monto_list_raw, str) and "," in monto_list_raw:
                            monto_list = [float(m.strip()) if m.strip() else 0.0 for m in monto_list_raw.split(",")]
                        else:
                            try:
                                monto_list = [float(monto_list_raw)] if monto_list_raw else []
                            except Exception:
                                monto_list = []

                        while len(fecha_list) < num_comprobantes:
                            fecha_list.append("")
                        while len(forma_list) < num_comprobantes:
                            forma_list.append("Transferencia")
                        while len(banco_list) < num_comprobantes:
                            banco_list.append("")
                        while len(terminal_list) < num_comprobantes:
                            terminal_list.append("")
                        while len(monto_list) < num_comprobantes:
                            monto_list.append(0.0)
                        while len(ref_list) < num_comprobantes:
                            ref_list.append("")

                        for i in range(num_comprobantes):
                            if num_comprobantes == 1:
                                st.markdown("### 🧾 Comprobante")
                            else:
                                emoji_num = ["1️⃣","2️⃣","3️⃣","4️⃣","5️⃣","6️⃣","7️⃣","8️⃣","9️⃣","🔟"]
                                label = emoji_num[i] if i < len(emoji_num) else str(i + 1)
                                st.markdown(f"### {label} 🧾 Comprobante {i+1}")

                            col_pago = st.columns(4)
                            with col_pago[0]:
                                fecha_raw = fecha_list[i] if i < len(fecha_list) else ""
                                parsed_fecha = None
                                if fecha_raw:
                                    parsed_fecha = pd.to_datetime(fecha_raw, errors="coerce")

                                if parsed_fecha is not None and not pd.isna(parsed_fecha):
                                    fecha_value = parsed_fecha.date()
                                else:
                                    fecha_value = None

                                fecha_i = st.date_input(
                                    f"📅 Fecha Pago {i+1}",
                                    value=fecha_value,
                                    key=f"fecha_pago_{i}"
                                )
                            with col_pago[1]:
                                forma_i = st.selectbox(
                                    f"💳 Forma de Pago {i+1}",
                                    ["Transferencia", "Depósito en Efectivo", "Tarjeta de Débito", "Tarjeta de Crédito", "Cheque"],
                                    index=["Transferencia", "Depósito en Efectivo", "Tarjeta de Débito", "Tarjeta de Crédito", "Cheque"].index(forma_list[i]) if forma_list[i] in ["Transferencia", "Depósito en Efectivo", "Tarjeta de Débito", "Tarjeta de Crédito", "Cheque"] else 0,
                                    key=f"forma_pago_{i}"
                                )
                            with col_pago[2]:
                                if forma_i in ["Tarjeta de Débito", "Tarjeta de Crédito"]:
                                    terminal_options = ["BANORTE", "AFIRME", "VELPAY", "CLIP", "PAYPAL", "BBVA", "CONEKTA", "MERCADO PAGO"]
                                    terminal_i = st.selectbox(
                                        f"🏧 Terminal {i+1}",
                                        terminal_options,
                                        index=terminal_options.index(terminal_list[i]) if terminal_list[i] in terminal_options else 0,
                                        key=f"terminal_pago_{i}"
                                    )
                                    banco_i = ""
                                else:
                                    banco_options = ["BANORTE", "BANAMEX", "AFIRME", "BANCOMER OP", "BANCOMER CURSOS"]
                                    banco_i = st.selectbox(
                                        f"🏦 Banco Destino {i+1}",
                                        banco_options,
                                        index=banco_options.index(banco_list[i]) if banco_list[i] in banco_options else 0,
                                        key=f"banco_pago_{i}"
                                    )
                                    terminal_i = ""

                            with col_pago[3]:
                                monto_i = st.number_input(
                                    f"💲 Monto {i+1}",
                                    min_value=0.0,
                                    format="%.2f",
                                    value=monto_list[i] if i < len(monto_list) else 0.0,
                                    key=f"monto_pago_{i}"
                                )

                            referencia_i = st.text_input(
                                f"🔢 Referencia {i+1}",
                                value=ref_list[i] if i < len(ref_list) else "",
                                key=f"ref_pago_{i}"
                            )

                            fecha_list[i] = fecha_i.strftime("%Y-%m-%d") if fecha_i else ""
                            forma_list[i] = forma_i
                            banco_list[i] = banco_i
                            terminal_list[i] = terminal_i
                            monto_list[i] = monto_i
                            ref_list[i] = referencia_i

                    reject_toggle_key = f"show_rechazo__{current_selection_key}"
                    reject_reason_key = f"motivo_rechazo__{current_selection_key}"
                    cancel_toggle_key = f"show_cancel__{current_selection_key}"
                    cancel_reason_key = f"motivo_cancelacion__{current_selection_key}"

                    col1, col2, col3, col4 = st.columns([2, 1, 1, 1])
                    with col1:
                        st.info("👆 Revisa los comprobantes antes de confirmar.")

                    confirm_disabled = num_comprobantes == 0 and not pedido_pagado_sin_confirmar
                    with col2:
                        if st.button(
                            "✅ Confirmar Comprobante",
                            use_container_width=True,
                            disabled=confirm_disabled,
                        ) and not confirm_disabled:
                            try:
                                gsheet_row_index = (
                                    df_pedidos[
                                        df_pedidos['ID_Pedido'] == selected_pedido_id_for_s3_search
                                    ].index[0]
                                    + 2
                                )

                                if FECHA_CONFIRMADO_COL not in df_pedidos.columns:
                                    df_pedidos[FECHA_CONFIRMADO_COL] = ""
                                fecha_confirmado = obtener_fecha_confirmado_cdmx()

                                updates = {
                                    'Comprobante_Confirmado': 'Sí',
                                    FECHA_CONFIRMADO_COL: fecha_confirmado,
                                    'Fecha_Pago_Comprobante': " y ".join(fecha_list),
                                    'Forma_Pago_Comprobante': ", ".join(forma_list),
                                    'Monto_Comprobante': sum(monto_list),
                                    'Referencia_Comprobante': ", ".join(ref_list),
                                    'Terminal': ", ".join([t for t in terminal_list if t]),
                                    'Banco_Destino_Pago': ", ".join([b for b in banco_list if b]),
                                }
                                if is_pedido_local:
                                    updates[ESTADO_ENTREGA_COL] = estado_entrega_value

                                # 🔹 OBTENER HOJA FRESCA (con reintentos) ANTES DE ESCRIBIR
                                worksheet = _get_ws_datos()
                                headers = ensure_sheet_column(worksheet, headers, FECHA_CONFIRMADO_COL)
                                if is_pedido_local:
                                    headers = ensure_sheet_column(worksheet, headers, ESTADO_ENTREGA_COL)
                                st.session_state.headers = headers

                                cell_updates = []
                                for col, val in updates.items():
                                    if col in headers:
                                        cell_updates.append({
                                            "range": rowcol_to_a1(
                                                gsheet_row_index, headers.index(col) + 1
                                            ),
                                            "values": [[val]],
                                        })

                                if cell_updates:
                                    safe_batch_update(worksheet, cell_updates)
                                    _get_ws_datos.clear()

                                df_idx = df_pedidos[
                                    df_pedidos['ID_Pedido'] == selected_pedido_id_for_s3_search
                                ].index
                                if len(df_idx) > 0:
                                    df_idx = df_idx[0]
                                    for col, val in updates.items():
                                        if col in df_pedidos.columns:
                                            df_pedidos.at[df_idx, col] = val
                                    st.session_state.df_pedidos = df_pedidos

                                pedidos_pagados_no_confirmados = refresh_pedidos_pagados_no_confirmados(
                                    df_pedidos
                                )
                                clear_comprobante_form_state()
                                st.session_state.pop("last_selected_pedido_key", None)
                                st.success("🎉 Comprobante confirmado exitosamente.")
                                st.balloons()
                                st.session_state["show_accept_after_confirm"] = True

                            except Exception as e:
                                st.error(f"❌ Error al confirmar comprobante: {e}")

                        if st.session_state.get("show_accept_after_confirm"):
                            if st.button(
                                "Aceptar",
                                key="confirm_accept_refresh",
                                use_container_width=True,
                            ):
                                st.session_state.pop("show_accept_after_confirm", None)
                                rerun_current_tab()

                    with col3:
                        if num_comprobantes == 0:
                            st.button(
                                "❌ Rechazar Comprobante",
                                use_container_width=True,
                                disabled=True,
                            )
                        else:
                            if st.button("❌ Rechazar Comprobante", use_container_width=True):
                                st.session_state[reject_toggle_key] = True
                                st.session_state.pop(cancel_toggle_key, None)
                            if st.session_state.get(reject_toggle_key):
                                st.session_state.setdefault(reject_reason_key, "")
                                st.text_area(
                                    "📝 Motivo de rechazo",
                                    key=reject_reason_key,
                                    placeholder="Describe el motivo del rechazo",
                                )
                                if st.button(
                                    "Guardar rechazo",
                                    key=f"guardar_rechazo_{current_selection_key}",
                                    use_container_width=True,
                                ):
                                    motivo = str(st.session_state.get(reject_reason_key, "")).strip()
                                    if not motivo:
                                        st.error("⚠️ Debes ingresar un motivo para rechazar el comprobante.")
                                    else:
                                        prefijo = f"Rechazo[{motivo}]"
                                        try:
                                            gsheet_row_index = (
                                                df_pedidos[
                                                    df_pedidos['ID_Pedido'] == selected_pedido_id_for_s3_search
                                                ].index[0]
                                                + 2
                                            )

                                            if FECHA_CONFIRMADO_COL not in df_pedidos.columns:
                                                df_pedidos[FECHA_CONFIRMADO_COL] = ""

                                            updates = {
                                                'Estado_Pago': '🔴 No Pagado',
                                                MOTIVO_RECHAZO_CANCELACION_COL: prefijo,
                                                'Comprobante_Confirmado': 'No',
                                                FECHA_CONFIRMADO_COL: '',
                                            }

                                            worksheet = _get_ws_datos()
                                            headers = ensure_sheet_column(worksheet, headers, FECHA_CONFIRMADO_COL)
                                            st.session_state.headers = headers

                                            cell_updates = []
                                            for col, val in updates.items():
                                                if col in headers:
                                                    cell_updates.append({
                                                        "range": rowcol_to_a1(
                                                            gsheet_row_index, headers.index(col) + 1
                                                        ),
                                                        "values": [[val]],
                                                    })

                                            if cell_updates:
                                                safe_batch_update(worksheet, cell_updates)
                                                _get_ws_datos.clear()
                                                cargar_pedidos_desde_google_sheet.clear()
                                            else:
                                                st.warning(
                                                    "⚠️ No se encontró la columna de motivos en la hoja. Verifica la configuración."
                                                )

                                            df_idx = df_pedidos[
                                                df_pedidos['ID_Pedido'] == selected_pedido_id_for_s3_search
                                            ].index
                                            if len(df_idx) > 0:
                                                df_idx = df_idx[0]
                                                for col, val in updates.items():
                                                    if col in df_pedidos.columns:
                                                        df_pedidos.at[df_idx, col] = val
                                                st.session_state.df_pedidos = df_pedidos

                                            pedidos_pagados_no_confirmados = refresh_pedidos_pagados_no_confirmados(
                                                df_pedidos
                                            )

                                            st.success("🚫 Comprobante rechazado correctamente.")
                                            st.session_state.pop(reject_toggle_key, None)
                                            st.session_state.pop(reject_reason_key, None)
                                            clear_comprobante_form_state()
                                            rerun_current_tab()

                                        except Exception as e:
                                            st.error(f"❌ Error al rechazar el comprobante: {e}")

                    with col4:
                        if st.button("🛑 Cancelar Pedido", use_container_width=True):
                            st.session_state[cancel_toggle_key] = True
                            st.session_state.pop(reject_toggle_key, None)
                        if st.session_state.get(cancel_toggle_key):
                            st.session_state.setdefault(cancel_reason_key, "")
                            st.text_area(
                                "📝 Motivo de cancelación",
                                key=cancel_reason_key,
                                placeholder="Describe el motivo de cancelación",
                            )
                            if st.button(
                                "Guardar cancelación",
                                key=f"guardar_cancelacion_{current_selection_key}",
                                use_container_width=True,
                            ):
                                motivo = str(st.session_state.get(cancel_reason_key, "")).strip()
                                if not motivo:
                                    st.error("⚠️ Debes ingresar un motivo para cancelar el pedido.")
                                else:
                                    prefijo = f"Cancelado[{motivo}]"
                                    try:
                                        gsheet_row_index = (
                                            df_pedidos[
                                                df_pedidos['ID_Pedido'] == selected_pedido_id_for_s3_search
                                            ].index[0]
                                            + 2
                                        )

                                        updates = {
                                            MOTIVO_RECHAZO_CANCELACION_COL: prefijo,
                                        }

                                        worksheet = _get_ws_datos()

                                        cell_updates = []
                                        for col, val in updates.items():
                                            if col in headers:
                                                cell_updates.append({
                                                    "range": rowcol_to_a1(
                                                        gsheet_row_index, headers.index(col) + 1
                                                    ),
                                                    "values": [[val]],
                                                })

                                        if cell_updates:
                                            safe_batch_update(worksheet, cell_updates)
                                            _get_ws_datos.clear()
                                            cargar_pedidos_desde_google_sheet.clear()
                                        else:
                                            st.warning(
                                                "⚠️ No se encontró la columna de motivos en la hoja. Verifica la configuración."
                                            )

                                        df_idx = df_pedidos[
                                            df_pedidos['ID_Pedido'] == selected_pedido_id_for_s3_search
                                        ].index
                                        if len(df_idx) > 0:
                                            df_idx = df_idx[0]
                                            for col, val in updates.items():
                                                if col in df_pedidos.columns:
                                                    df_pedidos.at[df_idx, col] = val
                                            st.session_state.df_pedidos = df_pedidos

                                        pedidos_pagados_no_confirmados = refresh_pedidos_pagados_no_confirmados(
                                            df_pedidos
                                        )

                                        st.success("🛑 Pedido cancelado y ocultado de la vista.")
                                        st.session_state.pop(cancel_toggle_key, None)
                                        st.session_state.pop(cancel_reason_key, None)
                                        clear_comprobante_form_state()
                                        rerun_current_tab()

                                    except Exception as e:
                                        st.error(f"❌ Error al cancelar el pedido: {e}")
# --- TAB 2: PEDIDOS CONFIRMADOS ---
with tab2:
    if tab2_is_active:
        st.session_state[TAB_SESSION_KEY] = 1
        st.session_state["current_tab"] = "1"
    st.header("📥 Pedidos Confirmados")

    # Imports usados en este bloque
    from io import BytesIO
    from datetime import datetime
    import gspread
    import re

    # Asegura nonce para esta pestaña
    if "tab2_reload_nonce" not in st.session_state:
        st.session_state["tab2_reload_nonce"] = 0

    # ✅ Cache de lectura (robusta, con snapshot)
    def dedupe_confirmados(df: pd.DataFrame) -> pd.DataFrame:
        """Normaliza y deduplica confirmados usando ID_Pedido y Folio_Factura."""
        if df is None or not isinstance(df, pd.DataFrame):
            return df

        if df.empty:
            return df.copy()

        trabajo = df.copy()

        if "ID_Pedido" not in trabajo.columns:
            trabajo["ID_Pedido"] = ""
        if "Folio_Factura" not in trabajo.columns:
            trabajo["Folio_Factura"] = ""

        trabajo["ID_Pedido"] = trabajo["ID_Pedido"].apply(normalize_id_pedido)
        trabajo["Folio_Factura"] = trabajo["Folio_Factura"].apply(normalize_folio_factura)
        trabajo["Folio_Factura"] = trabajo["Folio_Factura"].fillna("")

        trabajo["__folio_key"] = trabajo["Folio_Factura"].replace("", "__EMPTY__")
        trabajo["__dedupe_order"] = list(range(len(trabajo)))

        deduplicado = (
            trabajo.sort_values("__dedupe_order")
            .groupby(["ID_Pedido", "__folio_key"], as_index=False, sort=False)
            .tail(1)
        )

        deduplicado = deduplicado.sort_values("__dedupe_order")
        deduplicado = deduplicado[deduplicado["ID_Pedido"] != ""]
        deduplicado = deduplicado.drop(columns=["__folio_key", "__dedupe_order"], errors="ignore")
        deduplicado.reset_index(drop=True, inplace=True)

        return deduplicado

    def sync_estado_surtido_confirmados(
        df: pd.DataFrame, df_pedidos_src: pd.DataFrame
    ) -> pd.DataFrame:
        """Actualiza columnas clave en confirmados usando los últimos datos de pedidos."""
        if df is None or not isinstance(df, pd.DataFrame):
            return df

        if df.empty:
            return df.copy()

        if df_pedidos_src is None or not isinstance(df_pedidos_src, pd.DataFrame):
            return df.copy()

        if df_pedidos_src.empty:
            return df.copy()

        if "ID_Pedido" not in df.columns:
            return df.copy()

        if "ID_Pedido" not in df_pedidos_src.columns:
            return df.copy()

        trabajo = df.copy()

        pedidos_norm = df_pedidos_src.copy()
        pedidos_norm["__norm_id"] = pedidos_norm["ID_Pedido"].apply(normalize_id_pedido)
        pedidos_norm = pedidos_norm[pedidos_norm["__norm_id"] != ""]

        if pedidos_norm.empty:
            return trabajo

        pedidos_norm = pedidos_norm.drop_duplicates(subset="__norm_id", keep="last")
        trabajo["__norm_id"] = trabajo["ID_Pedido"].apply(normalize_id_pedido)

        pedidos_norm = pedidos_norm.set_index("__norm_id")

        columnas_disponibles = {
            destino: origen
            for destino, origen in CONFIRMADOS_SYNC_COLUMN_MAP.items()
            if origen in pedidos_norm.columns
        }

        if not columnas_disponibles:
            trabajo.drop(columns="__norm_id", inplace=True, errors="ignore")
            return trabajo

        for col_destino, col_origen in columnas_disponibles.items():
            if col_destino not in trabajo.columns:
                trabajo[col_destino] = pd.NA

            serie_origen = pedidos_norm[col_origen]
            mapped = trabajo["__norm_id"].map(serie_origen)
            mask_update = mapped.notna()

            if mask_update.any():
                trabajo.loc[mask_update, col_destino] = mapped.loc[mask_update]

        trabajo.drop(columns="__norm_id", inplace=True, errors="ignore")

        return trabajo

    @st.cache_data(show_spinner=False)
    def cargar_confirmados_guardados_cached(sheet_id: str, ws_name: str, _nonce: int):
        """
        Lee la hoja de confirmados con reintentos (safe_open_worksheet) y guarda snapshot.
        _nonce fuerza recarga manual.
        """
        ws = safe_open_worksheet(sheet_id, ws_name)
        vals = ws.get_values("A1:ZZ", value_render_option="UNFORMATTED_VALUE")
        if not vals:
            return pd.DataFrame(), [], 0, 0
        headers = vals[0]
        df = pd.DataFrame(vals[1:], columns=headers)

        total_filas_hoja = len(df)

        def _normalizar_columna_fecha(df_in: pd.DataFrame, columna: str) -> None:
            if columna not in df_in.columns:
                return

            def _convertir(valor):
                if isinstance(valor, pd.Timestamp):
                    return valor.strftime("%Y-%m-%d")
                if isinstance(valor, numbers.Number) and not isinstance(valor, bool) and not pd.isna(valor):
                    fecha = pd.to_datetime(valor, unit="D", origin="1899-12-30", errors="coerce")
                    if pd.isna(fecha):
                        return ""
                    return fecha.strftime("%Y-%m-%d")
                if valor is None or pd.isna(valor):
                    return ""
                return str(valor)

            df_in[columna] = df_in[columna].apply(_convertir)

        # Normalización mínima / columnas clave
        campos_clave = ['ID_Pedido', 'Cliente', 'Folio_Factura']
        for c in campos_clave:
            if c not in df.columns:
                df[c] = ""
        df = df.dropna(how='all')
        df = df[df[campos_clave].apply(
            lambda row: any(str(val).strip().lower() not in ["", "nan", "n/a"] for val in row),
            axis=1
        )]

        registros_antes_deduplicado = len(df)

        if not df.empty:
            df["__sheet_row"] = df.index + 2
            df = dedupe_confirmados(df)
            for columna_fecha in ("Fecha_Pago_Comprobante", "Fecha_Entrega"):
                _normalizar_columna_fecha(df, columna_fecha)
            if "Hora_Registro" in df.columns:
                def _formatear_hora_registro(valor):
                    if isinstance(valor, pd.Timestamp):
                        return valor.strftime("%d/%m/%Y %H:%M")
                    if (
                        isinstance(valor, numbers.Number)
                        and not isinstance(valor, bool)
                        and not pd.isna(valor)
                    ):
                        ts = pd.to_datetime(
                            valor,
                            unit="D",
                            origin="1899-12-30",
                            errors="coerce",
                        )
                        if pd.isna(ts):
                            return valor
                        return ts.strftime("%d/%m/%Y %H:%M")
                    if valor is None or pd.isna(valor):
                        return ""
                    return valor if isinstance(valor, str) else str(valor)

                df["Hora_Registro"] = df["Hora_Registro"].apply(_formatear_hora_registro)
        deduplicados = max(registros_antes_deduplicado - len(df), 0)

        # Snapshot "último bueno"
        st.session_state["_lastgood_confirmados"] = (df.copy(), headers[:])
        return df, headers, deduplicados, total_filas_hoja

    # 📄 Cargar hoja 'pedidos_confirmados' con fallback a snapshot si la API falla
    try:
        df_confirmados_guardados, headers_confirmados, duplicados_eliminados, total_original = cargar_confirmados_guardados_cached(
            GOOGLE_SHEET_ID, "pedidos_confirmados", st.session_state["tab2_reload_nonce"]
        )
    except gspread.exceptions.WorksheetNotFound:
        spreadsheet = get_spreadsheet(GOOGLE_SHEET_ID)
        spreadsheet.add_worksheet(title="pedidos_confirmados", rows=1000, cols=30)
        df_confirmados_guardados, headers_confirmados = pd.DataFrame(), []
        duplicados_eliminados = 0
        total_original = 0
    except gspread.exceptions.APIError as e:
        snap = st.session_state.get("_lastgood_confirmados")
        if snap:
            st.warning("♻️ Error temporal al leer 'pedidos_confirmados'. Mostrando último snapshot bueno.")
            df_confirmados_guardados, headers_confirmados = snap
            duplicados_eliminados = st.session_state.get("_last_confirmados_dedup", 0)
            total_original = st.session_state.get("_last_confirmados_total", len(df_confirmados_guardados))
        else:
            st.error(f"❌ No se pudo leer 'pedidos_confirmados'. Detalle: {e}")
            df_confirmados_guardados, headers_confirmados = pd.DataFrame(), []
            duplicados_eliminados = 0
            total_original = 0

    if (
        isinstance(df_confirmados_guardados, pd.DataFrame)
        and "Estado_Surtido_Almacen" not in df_confirmados_guardados.columns
    ):
        alias_estado = None
        for posible in ("Estado", "Estafo"):
            if posible in df_confirmados_guardados.columns:
                alias_estado = posible
                break

        if alias_estado:
            df_confirmados_guardados = df_confirmados_guardados.rename(
                columns={alias_estado: "Estado_Surtido_Almacen"}
            )
            headers_confirmados = [
                "Estado_Surtido_Almacen" if h == alias_estado else h
                for h in headers_confirmados
            ]

    df_confirmados_guardados = sync_estado_surtido_confirmados(
        df_confirmados_guardados, df_pedidos
    )

    if isinstance(df_confirmados_guardados, pd.DataFrame):
        st.session_state["_lastgood_confirmados"] = (
            df_confirmados_guardados.copy(),
            headers_confirmados[:],
        )

    st.session_state["_last_confirmados_dedup"] = duplicados_eliminados
    st.session_state["_last_confirmados_total"] = total_original

    if duplicados_eliminados > 0 and not df_confirmados_guardados.empty:
        st.info(
            f"ℹ️ Se detectaron {duplicados_eliminados} filas duplicadas en la hoja. La vista y las descargas ya están depuradas;"
            " al actualizar se limpiará la hoja en Drive."
        )
        if total_original:
            st.caption(
                f"Se muestran {len(df_confirmados_guardados)} confirmados únicos de {total_original} filas originales."
            )
        st.session_state["_confirmados_cleanup_snapshot"] = {
            "headers": headers_confirmados[:],
            "df": df_confirmados_guardados.copy(),
            "duplicados": duplicados_eliminados,
        }
    else:
        st.session_state.pop("_confirmados_cleanup_snapshot", None)

    # Métricas rápidas (usa df_pedidos en memoria si existe)
    if ('df_pedidos' in locals() or 'df_pedidos' in globals()) and not df_pedidos.empty:
        col1, col2, col3, col4 = st.columns(4)
        with col1:
            st.metric("Total Pedidos", len(df_pedidos))
        with col2:
            pagados = df_pedidos[df_pedidos.get('Estado_Pago') == '✅ Pagado'] if 'Estado_Pago' in df_pedidos.columns else pd.DataFrame()
            st.metric("Pedidos Pagados", len(pagados))
        with col3:
            st.metric("Comprobantes Confirmados", len(df_confirmados_guardados))
        with col4:
            pendientes = len(pedidos_pagados_no_confirmados) if 'pedidos_pagados_no_confirmados' in locals() else 0
            st.metric("Pendientes Confirmación", pendientes)

    st.markdown("---")

    # 🔁 Botón único: Actualizar Enlaces (agregar nuevos) + Recargar tabla
    tab2_alert = st.empty()
    if st.button(
        "🔁 Actualizar Enlaces y Recargar Confirmados",
        type="primary",
        help="Agrega confirmados nuevos con enlaces y refresca la tabla",
    ):
        if allow_refresh("tab2_last_refresh", tab2_alert):
            try:
                # Detectar nuevos confirmados no guardados aún en la hoja
                cleanup_snapshot = st.session_state.get("_confirmados_cleanup_snapshot")
                if cleanup_snapshot:
                    try:
                        spreadsheet = get_spreadsheet(GOOGLE_SHEET_ID)
                        try:
                            hoja_confirmados = spreadsheet.worksheet("pedidos_confirmados")
                        except gspread.exceptions.WorksheetNotFound:
                            hoja_confirmados = spreadsheet.add_worksheet(title="pedidos_confirmados", rows=1000, cols=30)

                        df_saneado = cleanup_snapshot["df"].copy()
                        df_saneado = df_saneado.drop(columns=["__sheet_row"], errors="ignore")
                        headers_saneados = cleanup_snapshot["headers"]
                        columnas_orden = headers_saneados if headers_saneados else df_saneado.columns.tolist()

                        for col in columnas_orden:
                            if col not in df_saneado.columns:
                                df_saneado[col] = ""
                        columnas_finales = columnas_orden[:]
                        for col in df_saneado.columns:
                            if col not in columnas_finales:
                                columnas_finales.append(col)

                        df_saneado = df_saneado[columnas_finales].fillna("").astype(str)
                        valores_actualizados = [columnas_finales] + df_saneado.values.tolist()

                        hoja_confirmados.clear()
                        hoja_confirmados.update("A1", valores_actualizados, value_input_option="USER_ENTERED")

                        tab2_alert.success(
                            f"🧹 Se limpiaron {cleanup_snapshot['duplicados']} duplicados directamente en la hoja."
                        )
                        st.session_state.pop("_confirmados_cleanup_snapshot", None)
                        df_confirmados_guardados = df_saneado
                        headers_confirmados = columnas_finales
                    except Exception as e:
                        tab2_alert.warning(
                            f"⚠️ No fue posible limpiar los duplicados en la hoja: {e}. Se continuará con la actualización."
                        )

                # Detectar nuevos confirmados no guardados aún en la hoja
                if not df_confirmados_guardados.empty:
                    ids_existentes_norm = df_confirmados_guardados["ID_Pedido"].apply(normalize_id_pedido)
                    folios_existentes_norm = df_confirmados_guardados["Folio_Factura"].apply(normalize_folio_factura)
                    pares_existentes = {
                        (id_val, folio_val)
                        for id_val, folio_val in zip(ids_existentes_norm, folios_existentes_norm)
                        if id_val
                    }
                else:
                    pares_existentes = set()

                if 'ID_Pedido' in df_pedidos.columns:
                    serie_ids_normalizados = df_pedidos['ID_Pedido'].apply(normalize_id_pedido)
                else:
                    serie_ids_normalizados = pd.Series(["" for _ in range(len(df_pedidos))], index=df_pedidos.index)

                if 'Folio_Factura' in df_pedidos.columns:
                    serie_folios_normalizados = df_pedidos['Folio_Factura'].apply(normalize_folio_factura)
                else:
                    serie_folios_normalizados = pd.Series(["" for _ in range(len(df_pedidos))], index=df_pedidos.index)

                pares_normalizados = pd.Series(
                    list(zip(serie_ids_normalizados, serie_folios_normalizados)), index=df_pedidos.index
                )

                df_nuevos = df_pedidos[
                    (df_pedidos.get('Comprobante_Confirmado') == 'Sí') &
                    (~pares_normalizados.isin(pares_existentes))
                ].copy()

                if not df_nuevos.empty:
                    df_nuevos.loc[:, 'ID_Pedido'] = serie_ids_normalizados.loc[df_nuevos.index]
                    df_nuevos.loc[:, 'Folio_Factura'] = serie_folios_normalizados.loc[df_nuevos.index]
                    df_nuevos = df_nuevos[(df_nuevos['ID_Pedido'] != "")]

                columnas_objetivo_confirmados = COLUMNAS_OBJETIVO_CONFIRMADOS[:]

                nuevos_agregados = 0

                if df_nuevos.empty:
                    tab2_alert.info("✅ No hay pedidos confirmados nuevos por registrar. Se recargará la tabla igualmente…")
                else:
                    df_nuevos = df_nuevos.sort_values(by='Fecha_Pago_Comprobante', ascending=False, na_position='last')
                    df_nuevos = df_nuevos.drop_duplicates(
                        subset=['ID_Pedido', 'Folio_Factura'], keep='first'
                    )

                    for col in columnas_objetivo_confirmados:
                        if col not in df_nuevos.columns:
                            df_nuevos[col] = ""

                    (
                        adjuntos_map,
                        adjuntos_surtido_map,
                        adjuntos_guia_map,
                    ) = build_adjuntos_map_from_pedidos(df_pedidos)

                    fallback_columns = [
                        'Link_Adjuntos',
                        'Link_Adjuntos_Modificacion',
                        'Link_Adjuntos_Guia',
                        'Link_Refacturacion',
                    ]
                    links_fallback_map = build_link_fallback_map(df_pedidos, fallback_columns)

                    def _fallback_link(norm_id: str, column: str) -> str:
                        if not norm_id:
                            return ""
                        valores = links_fallback_map.get(norm_id)
                        if not valores:
                            return ""
                        return valores.get(column, "")

                    link_adjuntos, link_adjuntos_mod, link_adjuntos_guia, link_refacturaciones = [], [], [], []

                    for _, row in df_nuevos.iterrows():
                        pedido_id = row.get("ID_Pedido")
                        tipo_envio = row.get("Tipo_Envio")
                        normalized_id = normalize_id_pedido(pedido_id)

                        assets = discover_comprobante_assets(pedido_id, tipo_envio, s3_client)

                        link_adjuntos_value, _ = resolve_adjuntos_link(
                            pedido_id,
                            normalized_id,
                            adjuntos_map,
                            map_label="adjuntos",
                            category="adjuntos",
                            page_title="Adjuntos",
                            s3_client_instance=s3_client,
                        )
                        if not link_adjuntos_value:
                            link_adjuntos_value = assets.get("comprobante_link", "") or _fallback_link(
                                normalized_id, "Link_Adjuntos"
                            )

                        link_adjuntos_mod_value, _ = resolve_adjuntos_link(
                            pedido_id,
                            normalized_id,
                            adjuntos_surtido_map,
                            map_label="adjuntos_modificacion",
                            category="adjuntos-modificacion",
                            page_title="Adjuntos de Modificación",
                            s3_client_instance=s3_client,
                        )
                        if not link_adjuntos_mod_value:
                            link_adjuntos_mod_value = _fallback_link(
                                normalized_id, "Link_Adjuntos_Modificacion"
                            )

                        link_adjuntos_guia_value, _ = resolve_adjuntos_link(
                            pedido_id,
                            normalized_id,
                            adjuntos_guia_map,
                            map_label="adjuntos_guia",
                            category="adjuntos-guia",
                            page_title="Adjuntos de Guía",
                            s3_client_instance=s3_client,
                        )
                        if not link_adjuntos_guia_value:
                            link_adjuntos_guia_value = assets.get("guia_url", "") or _fallback_link(
                                normalized_id, "Link_Adjuntos_Guia"
                            )

                        link_refacturacion_value = assets.get("refact_url", "") or _fallback_link(
                            normalized_id, "Link_Refacturacion"
                        )

                        link_adjuntos.append(link_adjuntos_value)
                        link_adjuntos_mod.append(link_adjuntos_mod_value)
                        link_adjuntos_guia.append(link_adjuntos_guia_value)
                        link_refacturaciones.append(link_refacturacion_value)

                    df_nuevos["Link_Adjuntos"] = link_adjuntos
                    df_nuevos["Link_Adjuntos_Modificacion"] = link_adjuntos_mod
                    df_nuevos["Link_Adjuntos_Guia"] = link_adjuntos_guia
                    df_nuevos["Link_Refacturacion"] = link_refacturaciones

                    df_nuevos = df_nuevos.fillna("").astype(str)
                    df_nuevos = df_nuevos.reindex(columns=columnas_objetivo_confirmados, fill_value="")

                    df_existente_merge = df_confirmados_guardados.drop(columns=["__sheet_row"], errors="ignore")
                    df_existente_merge = dedupe_confirmados(df_existente_merge)

                    df_combined = pd.concat(
                        [df_existente_merge, df_nuevos],
                        ignore_index=True,
                        sort=False,
                    )
                    df_combined = dedupe_confirmados(df_combined)
                    df_combined = sync_estado_surtido_confirmados(
                        df_combined, df_pedidos
                    )
                    nuevos_agregados = max(len(df_combined) - len(df_existente_merge), 0)

                    for col in columnas_objetivo_confirmados:
                        if col not in df_combined.columns:
                            df_combined[col] = ""
                    columnas_a_descartar = [
                        col
                        for col in df_combined.columns
                        if col not in columnas_objetivo_confirmados and col != "__sheet_row"
                    ]
                    if columnas_a_descartar:
                        df_combined = df_combined.drop(columns=columnas_a_descartar, errors="ignore")

                    columnas_finales = columnas_objetivo_confirmados[:]

                    df_combined = df_combined.reindex(columns=columnas_finales, fill_value="")

                    df_guardar = df_combined.drop(columns=["__sheet_row"], errors="ignore")
                    df_guardar = df_guardar.reindex(columns=columnas_finales, fill_value="")
                    df_guardar = df_guardar.fillna("").astype(str)

                    valores_actualizados = [columnas_finales]
                    if not df_guardar.empty:
                        valores_actualizados += df_guardar.values.tolist()

                    spreadsheet = get_spreadsheet(GOOGLE_SHEET_ID)
                    try:
                        hoja_confirmados = spreadsheet.worksheet("pedidos_confirmados")
                    except gspread.exceptions.WorksheetNotFound:
                        hoja_confirmados = spreadsheet.add_worksheet(title="pedidos_confirmados", rows=1000, cols=30)

                    hoja_confirmados.clear()
                    hoja_confirmados.update("A1", valores_actualizados, value_input_option="USER_ENTERED")

                tab2_alert.success(
                    f"✅ {nuevos_agregados} nuevos pedidos confirmados agregados a la hoja."
                )

                # Recargar
                st.session_state["tab2_reload_nonce"] += 1
                cargar_confirmados_guardados_cached.clear()
                st.toast("Datos recargados", icon="🔄")
                rerun_current_tab()

            except gspread.exceptions.APIError as e:
                tab2_alert.error(f"❌ Error de Google API al actualizar/recargar: {e}")
            except Exception as e:
                tab2_alert.error(f"❌ Ocurrió un error al actualizar/recargar: {e}")

    # ---------- Vista de confirmados ----------
    if df_confirmados_guardados.empty:
        st.info("ℹ️ No hay registros en la hoja 'pedidos_confirmados'.")
    else:
        # 🔽 Ordenar para mostrar lo más reciente primero
        df_confirmados_guardados = df_confirmados_guardados.copy()

        def _to_dt(s):
            return pd.to_datetime(s, errors='coerce', dayfirst=True, infer_datetime_format=True)

        if "Fecha_Pago_Comprobante" in df_confirmados_guardados.columns:
            dt = _to_dt(df_confirmados_guardados["Fecha_Pago_Comprobante"])
            if dt.notna().any():
                df_confirmados_guardados = df_confirmados_guardados.assign(_dt=dt).sort_values("_dt", ascending=False, na_position='last').drop(columns="_dt")
            elif "Fecha_Entrega" in df_confirmados_guardados.columns:
                dt2 = _to_dt(df_confirmados_guardados["Fecha_Entrega"])
                df_confirmados_guardados = df_confirmados_guardados.assign(_dt=dt2).sort_values("_dt", ascending=False, na_position='last').drop(columns="_dt")
            else:
                df_confirmados_guardados = df_confirmados_guardados.iloc[::-1].reset_index(drop=True)
        elif "Fecha_Entrega" in df_confirmados_guardados.columns:
            dt2 = _to_dt(df_confirmados_guardados["Fecha_Entrega"])
            df_confirmados_guardados = df_confirmados_guardados.assign(_dt=dt2).sort_values("_dt", ascending=False, na_position='last').drop(columns="_dt")
        else:
            df_confirmados_guardados = df_confirmados_guardados.iloc[::-1].reset_index(drop=True)

        st.success(f"✅ {len(df_confirmados_guardados)} pedidos confirmados (últimos primero).")

        df_confirmados_visible = df_confirmados_guardados.drop(columns=["__sheet_row"], errors="ignore")
        columnas_ordenadas_visible = COLUMNAS_OBJETIVO_CONFIRMADOS + [
            col
            for col in df_confirmados_visible.columns
            if col not in COLUMNAS_OBJETIVO_CONFIRMADOS
        ]
        df_confirmados_visible = df_confirmados_visible.reindex(
            columns=columnas_ordenadas_visible, fill_value=""
        )

        df_confirmados_vista, columnas_expandidas_tabla = expand_link_adjuntos_columns(df_confirmados_visible)

        columnas_a_ocultar = list(dict.fromkeys([*columnas_expandidas_tabla]))
        columnas_a_ocultar_set = set(columnas_a_ocultar)
        columnas_prioritarias = [
            "ID_Pedido",
            "Hora_Registro",
            "Folio_Factura",
            "Folio_Factura_Refacturada",
            "Cliente",
            "Vendedor_Registro",
            "Tipo_Envio",
            "Estado_Surtido_Almacen",
            ESTADO_ENTREGA_COL,
            "Estado_Pago",
            "Comprobante_Confirmado",
            FECHA_CONFIRMADO_COL,
            "Refacturacion_Tipo",
            "Refacturacion_Subtipo",
            "Forma_Pago_Comprobante",
            "Monto_Comprobante",
            "Fecha_Pago_Comprobante",
            "Banco_Destino_Pago",
            "Terminal",
            "Referencia_Comprobante",
            "Modificacion_Surtido",
        ]

        columnas_para_tabla = [
            col
            for col in columnas_prioritarias
            if col in df_confirmados_vista.columns and col not in columnas_a_ocultar_set
        ]
        columnas_para_tabla.extend(
            [
                col
                for col in COLUMNAS_OBJETIVO_CONFIRMADOS
                if col not in columnas_para_tabla
                and col in df_confirmados_vista.columns
                and col not in columnas_a_ocultar_set
            ]
        )
        columnas_para_tabla.extend(
            [
                col
                for col in df_confirmados_vista.columns
                if col not in columnas_para_tabla and col not in columnas_a_ocultar_set
            ]
        )

        st.dataframe(
            df_confirmados_vista[columnas_para_tabla] if columnas_para_tabla else df_confirmados_vista,
            use_container_width=True, hide_index=True
        )

        # Descargar Excel (desde el DF ya ordenado)
        df_excel, columnas_expandidas_excel = expand_link_adjuntos_columns(df_confirmados_visible)
        columnas_a_ocultar_excel = list(dict.fromkeys([*columnas_expandidas_excel]))
        if columnas_a_ocultar_excel:
            df_excel = df_excel.drop(columnas_a_ocultar_excel, axis=1, errors="ignore")

        columnas_excel_orden = COLUMNAS_OBJETIVO_CONFIRMADOS + [
            col
            for col in df_excel.columns
            if col not in COLUMNAS_OBJETIVO_CONFIRMADOS
        ]
        df_excel = df_excel.reindex(columns=columnas_excel_orden, fill_value="")
        df_excel = df_excel.fillna("")

        output_confirmados = BytesIO()
        with pd.ExcelWriter(output_confirmados, engine='xlsxwriter') as writer:
            df_excel.to_excel(writer, index=False, sheet_name='Confirmados')
        data_xlsx = output_confirmados.getvalue()

        st.download_button(
            label="📥 Descargar Excel Confirmados (últimos primero)",
            data=data_xlsx,
            file_name=f"confirmados_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        st.markdown("---")
        st.subheader("🚚 Actualizar estado de entrega de pedidos locales")

        if "Tipo_Envio" in df_confirmados_guardados.columns:
            df_local_confirmados = df_confirmados_guardados[
                df_confirmados_guardados["Tipo_Envio"].astype(str).str.strip()
                == "📍 Pedido Local"
            ].copy()
        else:
            df_local_confirmados = pd.DataFrame()

        if ESTADO_ENTREGA_COL not in df_local_confirmados.columns:
            df_local_confirmados[ESTADO_ENTREGA_COL] = ESTADO_ENTREGA_DEFAULT

        if df_local_confirmados.empty:
            st.info("ℹ️ No hay pedidos locales confirmados para actualizar.")
        else:
            if "ID_Pedido" in df_local_confirmados.columns:
                df_local_confirmados = df_local_confirmados.drop_duplicates(
                    subset="ID_Pedido", keep="last"
                )

            fecha_sort_col = "_fecha_confirmado_sort"
            if FECHA_CONFIRMADO_COL in df_local_confirmados.columns:
                df_local_confirmados[fecha_sort_col] = pd.to_datetime(
                    df_local_confirmados[FECHA_CONFIRMADO_COL],
                    errors="coerce",
                )
            else:
                df_local_confirmados[fecha_sort_col] = pd.NaT

            df_local_confirmados = df_local_confirmados.sort_values(
                fecha_sort_col,
                ascending=False,
                na_position="last",
            )

            df_local_confirmados[ESTADO_ENTREGA_COL] = df_local_confirmados[
                ESTADO_ENTREGA_COL
            ].apply(normalize_estado_entrega)

            busqueda_local = st.text_input(
                "🔍 Buscar pedido por folio o cliente",
                key="buscar_pedido_local_entrega_confirmados",
            ).strip()

            if busqueda_local:
                criterio = busqueda_local.lower()

                folios = (
                    df_local_confirmados.get("Folio_Factura", pd.Series(dtype=str))
                    .astype(str)
                    .str.lower()
                )
                clientes = (
                    df_local_confirmados.get("Cliente", pd.Series(dtype=str))
                    .astype(str)
                    .str.lower()
                )

                mask_busqueda = (
                    folios.str.contains(criterio, na=False)
                    | clientes.str.contains(criterio, na=False)
                )
                df_local_filtrado = df_local_confirmados[mask_busqueda].copy()
            else:
                df_local_filtrado = df_local_confirmados.copy()

            df_local_filtrado = df_local_filtrado.drop(columns=[fecha_sort_col], errors="ignore")
            df_local_confirmados = df_local_confirmados.drop(
                columns=[fecha_sort_col], errors="ignore"
            )

            def _local_display_label(row: pd.Series) -> str:
                folio = row.get("Folio_Factura", "N/A")
                cliente = row.get("Cliente", "N/A")
                estado = row.get(ESTADO_ENTREGA_COL, "") or "Sin estado"
                return f"📄 {folio} - 👤 {cliente} · {estado}"

            if df_local_filtrado.empty:
                st.info("ℹ️ No se encontraron pedidos que coincidan con la búsqueda.")
            else:
                local_options = {
                    idx: _local_display_label(row)
                    for idx, row in df_local_filtrado.iterrows()
                }

                if not local_options:
                    st.info("ℹ️ No hay pedidos locales disponibles para actualizar.")
                else:
                    option_keys = list(local_options.keys())
                    selected_idx_local = st.selectbox(
                        "Selecciona un pedido local para ajustar su entrega",
                        options=option_keys,
                        format_func=lambda i: local_options[i],
                    key="select_local_entrega_confirmados",
                )

                selected_local_row = df_local_filtrado.loc[selected_idx_local]
                estado_actual = normalize_estado_entrega(
                    selected_local_row.get(ESTADO_ENTREGA_COL, "")
                )
                if estado_actual not in ESTADO_ENTREGA_OPCIONES:
                    estado_actual = ESTADO_ENTREGA_DEFAULT
                estado_nuevo = st.radio(
                    "Estado de entrega",
                    ESTADO_ENTREGA_OPCIONES,
                    index=ESTADO_ENTREGA_OPCIONES.index(estado_actual),
                    horizontal=True,
                    key=f"radio_estado_entrega_{selected_idx_local}",
                )

                if st.button(
                    "💾 Guardar estado de entrega",
                    type="primary",
                    key=f"btn_guardar_estado_entrega_{selected_idx_local}",
                ):
                    try:
                        pedido_id_raw = selected_local_row.get("ID_Pedido", "")
                        pedido_id_norm = normalize_id_pedido(pedido_id_raw)
                        folio_norm = normalize_folio_factura(
                            selected_local_row.get("Folio_Factura", "")
                        )

                        df_pedidos_idx = pd.Index([])
                        if "ID_Pedido" in df_pedidos.columns and pedido_id_norm:
                            df_pedidos_idx = df_pedidos[
                                df_pedidos["ID_Pedido"].apply(normalize_id_pedido)
                                == pedido_id_norm
                            ].index

                        if df_pedidos_idx.empty and "Folio_Factura" in df_pedidos.columns and folio_norm:
                            df_pedidos_idx = df_pedidos[
                                df_pedidos["Folio_Factura"].apply(normalize_folio_factura)
                                == folio_norm
                            ].index

                        if df_pedidos_idx.empty:
                            raise ValueError(
                                "No se encontró el pedido en la hoja 'datos_pedidos'."
                            )

                        gsheet_row_index = df_pedidos_idx[0] + 2

                        worksheet = _get_ws_datos()
                        headers_local = ensure_sheet_column(
                            worksheet, headers, ESTADO_ENTREGA_COL
                        )
                        st.session_state.headers = headers_local
                        headers = headers_local

                        safe_batch_update(
                            worksheet,
                            [
                                {
                                    "range": rowcol_to_a1(
                                        gsheet_row_index,
                                        headers_local.index(ESTADO_ENTREGA_COL) + 1,
                                    ),
                                    "values": [[estado_nuevo]],
                                }
                            ],
                        )

                        _get_ws_datos.clear()
                        cargar_pedidos_desde_google_sheet.clear()

                        df_pedidos.at[df_pedidos_idx[0], ESTADO_ENTREGA_COL] = estado_nuevo
                        st.session_state.df_pedidos = df_pedidos

                        st.success("✅ Estado de entrega actualizado correctamente.")
                        st.toast("Estado de entrega actualizado", icon="📦")
                        rerun_current_tab()
                    except Exception as err:
                        st.error(f"❌ No se pudo actualizar el estado de entrega: {err}")

        st.markdown("---")
        st.subheader("🚚 Actualizar estado de entrega de pedidos locales no pagados")

        if "Tipo_Envio" in df_pedidos.columns:
            df_local_no_pagados = df_pedidos[
                df_pedidos["Tipo_Envio"].astype(str).str.strip() == "📍 Pedido Local"
            ].copy()
        else:
            df_local_no_pagados = pd.DataFrame()

        if "Estado_Pago" in df_local_no_pagados.columns:
            df_local_no_pagados = df_local_no_pagados[
                df_local_no_pagados["Estado_Pago"].astype(str).str.strip()
                == "🔴 No Pagado"
            ].copy()

        if ESTADO_ENTREGA_COL not in df_local_no_pagados.columns:
            df_local_no_pagados[ESTADO_ENTREGA_COL] = ESTADO_ENTREGA_DEFAULT

        if df_local_no_pagados.empty:
            st.info("ℹ️ No hay pedidos locales no pagados para actualizar.")
        else:
            if "ID_Pedido" in df_local_no_pagados.columns:
                df_local_no_pagados = df_local_no_pagados.drop_duplicates(
                    subset="ID_Pedido", keep="last"
                )

            df_local_no_pagados[ESTADO_ENTREGA_COL] = df_local_no_pagados[
                ESTADO_ENTREGA_COL
            ].apply(normalize_estado_entrega)

            busqueda_no_pagado = st.text_input(
                "🔍 Buscar pedido no pagado por folio o cliente",
                key="buscar_pedido_local_entrega_no_pagados",
            ).strip()

            if busqueda_no_pagado:
                criterio = busqueda_no_pagado.lower()

                folios = (
                    df_local_no_pagados.get("Folio_Factura", pd.Series(dtype=str))
                    .astype(str)
                    .str.lower()
                )
                clientes = (
                    df_local_no_pagados.get("Cliente", pd.Series(dtype=str))
                    .astype(str)
                    .str.lower()
                )

                mask_busqueda = (
                    folios.str.contains(criterio, na=False)
                    | clientes.str.contains(criterio, na=False)
                )
                df_local_filtrado = df_local_no_pagados[mask_busqueda].copy()
            else:
                df_local_filtrado = df_local_no_pagados.copy()

            def _local_display_label_no_pagado(row: pd.Series) -> str:
                folio = row.get("Folio_Factura", "N/A")
                cliente = row.get("Cliente", "N/A")
                estado = row.get(ESTADO_ENTREGA_COL, "") or "Sin estado"
                return f"📄 {folio} - 👤 {cliente} · {estado}"

            if df_local_filtrado.empty:
                st.info("ℹ️ No se encontraron pedidos que coincidan con la búsqueda.")
            else:
                local_options = {
                    idx: _local_display_label_no_pagado(row)
                    for idx, row in df_local_filtrado.iterrows()
                }

                if not local_options:
                    st.info("ℹ️ No hay pedidos locales disponibles para actualizar.")
                else:
                    option_keys = list(local_options.keys())
                    selected_idx_local = st.selectbox(
                        "Selecciona un pedido local no pagado para ajustar su entrega",
                        options=option_keys,
                        format_func=lambda i: local_options[i],
                        key="select_local_entrega_no_pagados",
                    )

                selected_local_row = df_local_filtrado.loc[selected_idx_local]
                estado_actual = normalize_estado_entrega(
                    selected_local_row.get(ESTADO_ENTREGA_COL, "")
                )
                if estado_actual not in ESTADO_ENTREGA_OPCIONES:
                    estado_actual = ESTADO_ENTREGA_DEFAULT
                estado_nuevo = st.radio(
                    "Estado de entrega",
                    ESTADO_ENTREGA_OPCIONES,
                    index=ESTADO_ENTREGA_OPCIONES.index(estado_actual),
                    horizontal=True,
                    key=f"radio_estado_entrega_no_pagados_{selected_idx_local}",
                )

                if st.button(
                    "💾 Guardar estado de entrega",
                    type="primary",
                    key=f"btn_guardar_estado_entrega_no_pagados_{selected_idx_local}",
                ):
                    try:
                        pedido_id_raw = selected_local_row.get("ID_Pedido", "")
                        pedido_id_norm = normalize_id_pedido(pedido_id_raw)
                        folio_norm = normalize_folio_factura(
                            selected_local_row.get("Folio_Factura", "")
                        )

                        df_pedidos_idx = pd.Index([])
                        if "ID_Pedido" in df_pedidos.columns and pedido_id_norm:
                            df_pedidos_idx = df_pedidos[
                                df_pedidos["ID_Pedido"].apply(normalize_id_pedido)
                                == pedido_id_norm
                            ].index

                        if df_pedidos_idx.empty and "Folio_Factura" in df_pedidos.columns and folio_norm:
                            df_pedidos_idx = df_pedidos[
                                df_pedidos["Folio_Factura"].apply(normalize_folio_factura)
                                == folio_norm
                            ].index

                        if df_pedidos_idx.empty:
                            raise ValueError(
                                "No se encontró el pedido en la hoja 'datos_pedidos'."
                            )

                        gsheet_row_index = df_pedidos_idx[0] + 2

                        worksheet = _get_ws_datos()
                        headers_local = ensure_sheet_column(
                            worksheet, headers, ESTADO_ENTREGA_COL
                        )
                        st.session_state.headers = headers_local
                        headers = headers_local

                        safe_batch_update(
                            worksheet,
                            [
                                {
                                    "range": rowcol_to_a1(
                                        gsheet_row_index,
                                        headers_local.index(ESTADO_ENTREGA_COL) + 1,
                                    ),
                                    "values": [[estado_nuevo]],
                                }
                            ],
                        )

                        _get_ws_datos.clear()
                        cargar_pedidos_desde_google_sheet.clear()

                        df_pedidos.at[df_pedidos_idx[0], ESTADO_ENTREGA_COL] = estado_nuevo
                        st.session_state.df_pedidos = df_pedidos

                        st.success("✅ Estado de entrega actualizado correctamente.")
                        st.toast("Estado de entrega actualizado", icon="📦")
                        rerun_current_tab()
                    except Exception as err:
                        st.error(f"❌ No se pudo actualizar el estado de entrega: {err}")
                        
# --- TAB 3: CONFIRMACIÓN DE CASOS (Devoluciones + Garantías, con tabla y selectbox) ---
with tab3, suppress(StopException):
    if tab3_is_active:
        st.session_state[TAB_SESSION_KEY] = 2
        st.session_state["current_tab"] = "2"
    st.header("📦 Confirmación de Casos (Devoluciones + Garantías)")

    from datetime import datetime
    import uuid, os, json, math, re, time
    import pandas as pd
    import gspread
    # ⬇️ NUEVO: para generar/descargar el Word
    from docx import Document
    from io import BytesIO

    tab3_alert = st.empty()

    # ================== Helpers NUEVOS (solo para Word) ==================
    PLACEHOLDER_PATTERN = re.compile(r"\{\{([^}]+)\}\}")

    def _safe_value(v):
        if v is None:
            return "Sin registro"
        s = str(v).strip()
        return "Sin registro" if s.lower() in ("", "none", "nan", "n/a") else s

    def _iter_paragraphs_within(container):
        if container is None:
            return

        element = getattr(container, "_element", None)
        if element is not None:
            from docx.oxml.ns import qn
            from docx.text.paragraph import Paragraph

            parent = container
            for para_element in element.iter(qn("w:p")):
                yield Paragraph(para_element, parent)
            return

        for paragraph in getattr(container, "paragraphs", []):
            yield paragraph
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from _iter_paragraphs_within(cell)

    def _replace_span_in_runs(runs, start: int, end: int, replacement: str) -> None:
        current_pos = 0
        start_run_idx = None
        start_offset = 0
        end_run_idx = None
        end_offset = 0

        for idx, run in enumerate(runs):
            text = run.text or ""
            run_start = current_pos
            run_end = current_pos + len(text)

            if start_run_idx is None and start < run_end:
                start_run_idx = idx
                start_offset = start - run_start

            if start_run_idx is not None and end <= run_end:
                end_run_idx = idx
                end_offset = end - run_start
                break

            current_pos = run_end

        if start_run_idx is None or end_run_idx is None:
            return

        if start_run_idx == end_run_idx:
            run = runs[start_run_idx]
            text = run.text or ""
            run.text = text[:start_offset] + replacement + text[end_offset:]
            return

        start_run = runs[start_run_idx]
        end_run = runs[end_run_idx]

        start_text = start_run.text or ""
        end_text = end_run.text or ""

        start_run.text = start_text[:start_offset] + replacement
        end_run.text = end_text[end_offset:]

        for idx in range(start_run_idx + 1, end_run_idx):
            runs[idx].text = ""

    def _docx_replace_all(doc: Document, mapping: dict[str, str]) -> tuple[int, int, list[str]]:
        """Reemplaza placeholders {{...}} en todo el documento.

        Devuelve una tupla con (total_encontrados, total_reemplazados, placeholders_pendientes).
        """

        token_mapping = {f"{{{{{key}}}}}": _safe_value(value) for key, value in mapping.items()}
        total_found = 0
        total_replaced = 0

        def _process_paragraph(paragraph) -> None:
            nonlocal total_found, total_replaced

            runs = list(paragraph.runs)
            if not runs:
                text = paragraph.text or ""
                if not text:
                    return
                matches = PLACEHOLDER_PATTERN.findall(text)
                total_found += len(matches)
                new_text = text
                for token, replacement in token_mapping.items():
                    if token in new_text:
                        occurrences = new_text.count(token)
                        if occurrences:
                            new_text = new_text.replace(token, replacement)
                            total_replaced += occurrences
                if new_text != text:
                    paragraph.text = new_text
                return

            full_text = "".join(run.text or "" for run in runs)
            if not full_text:
                return

            matches = PLACEHOLDER_PATTERN.findall(full_text)
            total_found += len(matches)
            if not matches:
                return

            for token, replacement in token_mapping.items():
                if token == replacement:
                    continue
                token_len = len(token)
                current_text = "".join(run.text or "" for run in runs)
                if token not in current_text:
                    continue
                while True:
                    idx = current_text.find(token)
                    if idx == -1:
                        break
                    _replace_span_in_runs(runs, idx, idx + token_len, replacement)
                    total_replaced += 1
                    current_text = "".join(run.text or "" for run in runs)

        # Cuerpo principal
        for paragraph in _iter_paragraphs_within(doc):
            _process_paragraph(paragraph)

        # Encabezados y pies de página
        seen_parts: set[int] = set()
        for section in doc.sections:
            for attr in (
                "header",
                "first_page_header",
                "even_page_header",
                "footer",
                "first_page_footer",
                "even_page_footer",
            ):
                part = getattr(section, attr, None)
                part_id = id(part)
                if part is None or part_id in seen_parts:
                    continue
                seen_parts.add(part_id)
                for paragraph in _iter_paragraphs_within(part):
                    _process_paragraph(paragraph)

        # Revisión de placeholders restantes
        remaining = []
        for paragraph in _iter_paragraphs_within(doc):
            text = "".join((run.text or "") for run in paragraph.runs) or paragraph.text
            if not text:
                continue
            remaining.extend(PLACEHOLDER_PATTERN.findall(text))

        seen_parts.clear()
        for section in doc.sections:
            for attr in (
                "header",
                "first_page_header",
                "even_page_header",
                "footer",
                "first_page_footer",
                "even_page_footer",
            ):
                part = getattr(section, attr, None)
                part_id = id(part)
                if part is None or part_id in seen_parts:
                    continue
                seen_parts.add(part_id)
                for paragraph in _iter_paragraphs_within(part):
                    text = "".join((run.text or "") for run in paragraph.runs) or paragraph.text
                    if not text:
                        continue
                    remaining.extend(PLACEHOLDER_PATTERN.findall(text))

        return total_found, total_replaced, remaining
    # =====================================================================

    # Estado local
    if "tab3_reload_nonce" not in st.session_state:
        st.session_state["tab3_reload_nonce"] = 0
    if "tab3_selected_idx" not in st.session_state:
        st.session_state["tab3_selected_idx"] = 0

    # Lectura con fallback
    @st.cache_data(show_spinner=False)
    def get_raw_sheet_data_cached(sheet_id, worksheet_name, _nonce: int):
        try:
            try:
                ws = safe_open_worksheet(sheet_id, worksheet_name)
            except Exception:
                ws = get_spreadsheet(sheet_id).worksheet(worksheet_name)
            vals = ws.get_all_values()
            # guarda snapshot "último bueno" para futuros fallbacks
            st.session_state["_tab3_lastgood"] = vals
            return vals
        except gspread.exceptions.APIError:
            snap = st.session_state.get("_tab3_lastgood")
            if snap:
                return snap
            raise

    def process_sheet_data(raw_data):
        if not raw_data or len(raw_data) < 1:
            return pd.DataFrame(), []
        headers = raw_data[0]
        df = pd.DataFrame(raw_data[1:], columns=headers)
        return df, headers

    # Recargar
    col_recargar, _ = st.columns([1, 5])
    with col_recargar:
        def _reload_tab3():
            if not allow_refresh("tab3_last_refresh", tab3_alert):
                return
            st.session_state["tab3_reload_nonce"] += 1
            get_raw_sheet_data_cached.clear()
            st.toast("Casos recargados", icon="🔄")
            rerun_current_tab()

        if st.button(
            "🔄 Recargar casos",
            type="secondary",
            key="tab3_reload_btn",
        ):
            _reload_tab3()

    # Carga principal
    try:
        raw_casos = get_raw_sheet_data_cached(
            sheet_id=GOOGLE_SHEET_ID,
            worksheet_name="casos_especiales",
            _nonce=st.session_state["tab3_reload_nonce"],
        )
    except Exception as e:
        tab3_alert.error(f"❌ No se pudo leer 'casos_especiales'. {e}")
        st.stop()

    df_casos, headers_casos = process_sheet_data(raw_casos)
    if df_casos.empty:
        tab3_alert.info("ℹ️ No hay casos registrados en 'casos_especiales'.")
        st.stop()

    # Columnas mínimas
    needed_cols = [
        "ID_Pedido","Hora_Registro","Vendedor_Registro","Cliente","Folio_Factura",
        "Tipo_Envio","Resultado_Esperado","Numero_Cliente_RFC","Area_Responsable","Nombre_Responsable",
        "Material_Devuelto","Monto_Devuelto","Motivo_Detallado","Motivo_NotaVenta","Tipo_Envio_Original","Estatus_OrigenF",
        "Adjuntos","Hoja_Ruta_Mensajero","Estado_Caso","Estado_Recepcion","Turno","Fecha_Entrega",
        "Numero_Serie","Fecha_Compra","Seguimiento", GUIAS_DEVOLUCION_COL
    ]
    for c in needed_cols:
        if c not in df_casos.columns:
            df_casos[c] = ""

    # Utils
    def _is_blank(v: str) -> bool:
        s = str(v).strip().lower()
        return (s == "") or (s in ["nan","n/a","none"])

    def _norm(v: str) -> str:
        return str(v).strip()

    def pick_first_col(headers: list, candidates: list[str], default_if_missing: str = "") -> str:
        for c in candidates:
            if c in headers:
                return c
        if default_if_missing:
            headers.append(default_if_missing)
            return default_if_missing
        return ""

    def _normalize_urls(value):
        if value is None:
            return []
        if isinstance(value, float) and math.isnan(value):
            return []
        s = str(value).strip()
        if not s or s.lower() in ("nan","none","n/a"):
            return []
        urls = []
        try:
            obj = json.loads(s)
            if isinstance(obj, list):
                for it in obj:
                    if isinstance(it, str) and it.strip(): urls.append(it.strip())
                    elif isinstance(it, dict):
                        u = it.get("url") or it.get("URL") or it.get("href") or it.get("link")
                        if u and str(u).strip(): urls.append(str(u).strip())
            elif isinstance(obj, dict):
                for k in ("url","URL","link","href"):
                    if obj.get(k): urls.append(str(obj[k]).strip())
        except Exception:
            parts = re.split(r"[,\n;]+", s)
            for p in parts:
                p = p.strip()
                if p: urls.append(p)
        seen, out = set(), []
        for u in urls:
            if u not in seen:
                seen.add(u); out.append(u)
        return out

    def __s(v):
        return str(v).strip()

    def __has(v):
        s = __s(v)
        return bool(s) and s.lower() not in ("nan", "none", "n/a")

    def __is_url(v):
        s = __s(v).lower()
        return s.startswith("http://") or s.startswith("https://")

    def __link(url, label=None):
        u = __s(url)
        if __is_url(u):
            display_text = u
            safe_display = (
                display_text.replace("[", "\\[")
                .replace("]", "\\]")
                .replace("`", "\\`")
            )
            safe_url = (
                u.replace(" ", "%20")
                .replace("(", "%28")
                .replace(")", "%29")
            )
            return f"[{safe_display}]({safe_url})"
        return u

    def render_caso_especial(row):
        tipo = __s(row.get("Tipo_Envio", ""))
        is_dev = (tipo == "🔁 Devolución")
        title = "🧾 Caso Especial – 🔁 Devolución" if is_dev else "🧾 Caso Especial – 🛠 Garantía"
        st.markdown(f"### {title}")

        mod_txt = clean_modificacion_surtido(row.get("Modificacion_Surtido", ""))
        mod_txt_message = f"🛠 Modificación de surtido: {mod_txt}" if mod_txt else ""
        mod_txt_displayed = False

        if mod_txt_message:
            st.info(mod_txt_message)
            mod_txt_displayed = True

        vendedor = row.get("Vendedor_Registro", "") or row.get("Vendedor", "")
        hora = row.get("Hora_Registro", "")

        if is_dev:
            folio_nuevo = row.get("Folio_Factura", "")
            folio_error = row.get("Folio_Factura_Error", "")
            estatus_origen_raw = row.get("Estatus_OrigenF", "")
            estatus_origen_display = __s(estatus_origen_raw) if __has(estatus_origen_raw) else "N/A"
            st.markdown(
                f"📄 **Folio Nuevo:** `{folio_nuevo or 'N/A'}`  |  "
                f"📄 **Folio Error:** `{folio_error or 'N/A'}`  |  "
                f"📊 **Estatus Factura Origen:** `{estatus_origen_display}`  |  "
                f"🧑‍💼 **Vendedor:** `{vendedor or 'N/A'}`  |  "
                f"🕒 **Hora:** `{hora or 'N/A'}`"
            )
            guias_registradas = __s(row.get(GUIAS_DEVOLUCION_COL, ""))
            if __has(guias_registradas):
                st.markdown(f"**🚚 Número de guías registradas:** {guias_registradas}")
        else:
            st.markdown(
                f"📄 **Folio:** `{row.get('Folio_Factura','') or 'N/A'}`  |  "
                f"🧑‍💼 **Vendedor:** `{vendedor or 'N/A'}`  |  "
                f"🕒 **Hora:** `{hora or 'N/A'}`"
            )

        motivo_nota = __s(row.get("Motivo_NotaVenta", ""))
        if __has(motivo_nota):
            nota_val = __s(row.get("Folio_Factura", "")) or "N/A"
            st.markdown(f"**🧾 Nota de Venta:** `{nota_val}`")
            st.markdown("**✏️ Motivo de nota de venta:**")
            st.info(motivo_nota)

        st.markdown(
            f"**👤 Cliente:** {row.get('Cliente','N/A')}  |  **RFC:** {row.get('Numero_Cliente_RFC','') or 'N/A'}"
        )
        if not is_dev:
            st.markdown(
                f"**🔢 Número de Serie:** {row.get('Numero_Serie','') or 'N/A'}  |  "
                f"**🗓️ Fecha de Compra:** {row.get('Fecha_Compra','') or 'N/A'}"
            )
        st.markdown(
            f"**Estado:** {row.get('Estado','') or 'N/A'}  |  "
            f"**Estado del Caso:** {row.get('Estado_Caso','') or 'N/A'}  |  "
            f"**Turno:** {row.get('Turno','') or 'N/A'}"
        )

        rt = __s(row.get("Refacturacion_Tipo",""))
        rs = __s(row.get("Refacturacion_Subtipo",""))
        rf = __s(row.get("Folio_Factura_Refacturada",""))
        if __has(rt) or __has(rs) or __has(rf):
            st.markdown("**♻️ Refacturación:**")
            if __has(rt): st.markdown(f"- **Tipo:** {rt}")
            if __has(rs): st.markdown(f"- **Subtipo:** {rs}")
            if __has(rf): st.markdown(f"- **Folio refacturado:** {rf}")

        if __has(row.get("Resultado_Esperado","")):
            st.markdown(f"**🎯 Resultado Esperado:** {row.get('Resultado_Esperado')}")
        if __has(row.get("Motivo_Detallado","")):
            st.markdown("**📝 Motivo / Descripción:**")
            st.info(__s(row.get("Motivo_Detallado","")))
        if __has(row.get("Material_Devuelto","")):
            st.markdown("**📦 Piezas / Material:**")
            st.info(__s(row.get("Material_Devuelto","")))
        if __has(row.get("Monto_Devuelto","")):
            st.markdown(f"**💵 Monto (dev./estimado):** {row.get('Monto_Devuelto')}")

        if __has(row.get("Area_Responsable","")) or __has(row.get("Nombre_Responsable","")):
            st.markdown(
                f"**🏢 Área Responsable:** {row.get('Area_Responsable','') or 'N/A'}  |  "
                f"**👥 Responsable del Error:** {row.get('Nombre_Responsable','') or 'N/A'}"
            )

        if __has(row.get("Fecha_Entrega","")) or __has(row.get("Fecha_Recepcion_Devolucion","")) or __has(row.get("Estado_Recepcion","")):
            st.markdown(
                f"**📅 Fecha Entrega/Cierre:** {row.get('Fecha_Entrega','') or 'N/A'}  |  "
                f"**📅 Recepción:** {row.get('Fecha_Recepcion_Devolucion','') or 'N/A'}  |  "
                f"**📦 Recepción:** {row.get('Estado_Recepcion','') or 'N/A'}"
            )

        nota = __s(row.get("Nota_Credito_URL",""))
        docad = __s(row.get("Documento_Adicional_URL",""))
        if __has(nota):
            st.markdown(f"**🧾 Nota de Crédito:** {__link(nota, 'Nota de Crédito') if __is_url(nota) else nota}")
        if __has(docad):
            st.markdown(f"**📂 Documento Adicional:** {__link(docad, 'Documento Adicional') if __is_url(docad) else docad}")
        if __has(row.get("Comentarios_Admin_Devolucion","")):
            st.markdown("**🗒️ Comentario Administrativo:**")
            st.info(__s(row.get("Comentarios_Admin_Devolucion","")))

        adj_mod_raw = row.get("Adjuntos_Surtido","")
        if 'partir_urls' in globals():
            adj_mod = partir_urls(adj_mod_raw)
        else:
            adj_mod = [x.strip() for x in str(adj_mod_raw).split(",") if x.strip()]
        show_mod_section = bool(adj_mod) or (mod_txt_message and not mod_txt_displayed)
        if show_mod_section:
            st.markdown("#### 🛠 Modificación de surtido")
            if mod_txt_message and not mod_txt_displayed:
                st.info(mod_txt_message)
                mod_txt_displayed = True
            if adj_mod:
                st.markdown("**Archivos de modificación:**")
                for u in adj_mod:
                    st.markdown(f"- {__link(u)}")

        with st.expander("📎 Archivos (Adjuntos y Guía)", expanded=False):
            adj_raw = row.get("Adjuntos","")
            if 'partir_urls' in globals():
                adj = partir_urls(adj_raw)
            else:
                adj = [x.strip() for x in str(adj_raw).split(",") if x.strip()]
            guia = __s(row.get("Hoja_Ruta_Mensajero","")) or __s(row.get("Adjuntos_Guia",""))
            has_any = False
            if adj:
                has_any = True
                st.markdown("**Adjuntos:**")
                for u in adj:
                    st.markdown(f"- {__link(u)}")
            if __has(guia) and __is_url(guia):
                has_any = True
                st.markdown("**Guía:**")
                st.markdown(f"- {__link(guia, 'Abrir guía')}")
            if not has_any:
                st.info("Sin archivos registrados en la hoja.")

        st.markdown("---")

    # PENDIENTES (ambos tipos)
    mask_tipo_valido = df_casos["Tipo_Envio"].astype(str).str.strip().isin(["🔁 Devolución","🛠 Garantía"])
    estado_caso_norm = df_casos["Estado_Caso"].astype(str).apply(_norm).str.lower()
    mask_estado_caso_ok = (estado_caso_norm == "aprobado") | (estado_caso_norm == "")
    seguimiento_norm = df_casos["Seguimiento"].astype(str).apply(_norm).str.lower()
    mask_seguimiento_activo = (seguimiento_norm != "cerrado")

    df_pendientes = df_casos[mask_tipo_valido & mask_estado_caso_ok & mask_seguimiento_activo].copy()

    # ====== TABLA (se mantiene) ======
    if df_pendientes.empty:
        st.success("🎉 ¡No hay casos pendientes de confirmación!")
        st.stop()
    else:
        st.warning(f"📋 Hay {len(df_pendientes)} casos pendientes por confirmar.")

        columns_to_show = [
            "Tipo_Envio","Folio_Factura","Cliente","Vendedor_Registro",
            "Hora_Registro","Resultado_Esperado","Numero_Cliente_RFC",
            "Area_Responsable","Nombre_Responsable",
            "Material_Devuelto","Monto_Devuelto","Motivo_Detallado","Motivo_NotaVenta",
            "Tipo_Envio_Original","Estatus_OrigenF","Estado_Caso","Estado_Recepcion","Seguimiento"
        ]
        existing_columns = [c for c in columns_to_show if c in df_pendientes.columns]

        df_tabla = df_pendientes[existing_columns].copy()
        if "Hora_Registro" in df_tabla.columns:
            _hora = pd.to_datetime(df_tabla["Hora_Registro"], errors="coerce")
            df_tabla["Hora_Registro"] = _hora.dt.strftime("%d/%m/%Y %H:%M").fillna(df_tabla["Hora_Registro"])

        st.dataframe(
            df_tabla.sort_values(
                by="Hora_Registro" if "Hora_Registro" in df_tabla.columns else existing_columns[0],
                ascending=True
            ),
            use_container_width=True, hide_index=True
        )

    # ====== SELECTBOX (se mantiene, pero mezcla ambos tipos) ======
    df_pendientes["__Hora"] = pd.to_datetime(df_pendientes["Hora_Registro"], errors="coerce")
    df_pendientes = df_pendientes.sort_values(by="__Hora", ascending=True)

    # display incluye emoji del tipo
    df_pendientes["__display__"] = df_pendientes.apply(
        lambda rr: f"{rr.get('Tipo_Envio','')}  {str(rr.get('Folio_Factura','')).strip() or 's/folio'} – {str(rr.get('Cliente','')).strip() or 's/cliente'}  |  Esperado: {str(rr.get('Resultado_Esperado','')).strip()}",
        axis=1
    )
    options = df_pendientes["__display__"].tolist()

    if st.session_state["tab3_selected_idx"] >= len(options):
        st.session_state["tab3_selected_idx"] = 0

    selected = st.selectbox(
        "📋 Selecciona un caso",
        options,
        index=st.session_state["tab3_selected_idx"],
        key="tab3_selectbox",
    )
    st.session_state["tab3_selected_idx"] = options.index(selected) if selected in options else 0
    row = df_pendientes[df_pendientes["__display__"] == selected].iloc[0]

    # Índice real en hoja por ID
    matches = df_casos.index[df_casos["ID_Pedido"].astype(str).str.strip() == str(row["ID_Pedido"]).strip()]
    if len(matches) == 0:
        tab3_alert.error("❌ No se encontró el caso seleccionado en 'casos_especiales'.")
        st.stop()
    gsheet_row_idx = int(matches[0]) + 2

    # Worksheet para escritura con reintentos
    try:
        worksheet_casos = safe_open_worksheet(GOOGLE_SHEET_ID, "casos_especiales")
    except gspread.exceptions.APIError:
        tab3_alert.warning(
            "⚠️ Google Sheets está aplicando un cooldown. Se trabajará con el snapshot hasta que pase el bloqueo."
        )
        st.stop()

    # ========= RENDER DEL CASO SELECCIONADO (detecta si es Devolución o Garantía) =========
    tipo_case = str(row.get("Tipo_Envio","")).strip()
    is_dev = (tipo_case == "🔁 Devolución")
    # Garantía si no es devolución:
    is_gar = not is_dev

    render_caso_especial(row)

    # ===== FORMULARIO (ajusta columnas según tipo detectado) =====
    guias_key: str | None = None
    with st.form(key="tab3_confirm_form", clear_on_submit=False):
        fecha_key = f"fecha_recepcion_{'devolucion' if is_dev else 'garantia'}"
        estado_key = f"estado_recepcion_{'devolucion' if is_dev else 'garantia'}"
        seg_key = f"seguimiento_{row.get('ID_Pedido','')}"

        _fecha_raw = str(
            row.get(
                "Fecha_Recepcion_Devolucion" if is_dev else "Fecha_Recepcion_Garantia",
                "",
            )
        ).strip()
        _fecha_val = (
            pd.to_datetime(_fecha_raw, errors="coerce").date() if _fecha_raw else None
        )
        _estado_raw = str(row.get("Estado_Recepcion", "")).strip()
        _estado_val = (
            "Sí, completo" if _estado_raw == "Todo correcto" else (
                "Faltan artículos" if _estado_raw else None
            )
        )
        _segui_val = str(row.get("Seguimiento", "")).strip()

        if (
            fecha_key not in st.session_state
            or st.session_state.get("tab3_last_case_id") != row.get("ID_Pedido")
        ):
            st.session_state[fecha_key] = _fecha_val
        if (
            estado_key not in st.session_state
            or st.session_state.get("tab3_last_case_id") != row.get("ID_Pedido")
        ):
            st.session_state[estado_key] = _estado_val
        if (
            seg_key not in st.session_state
            or st.session_state.get("tab3_last_case_id") != row.get("ID_Pedido")
        ):
            st.session_state[seg_key] = _segui_val
        if is_dev:
            guias_key = f"guias_devolucion_{row.get('ID_Pedido','')}"
            _guias_raw = str(row.get(GUIAS_DEVOLUCION_COL, "")).strip()
            _guias_val = _guias_raw if _guias_raw in GUIAS_DEVOLUCION_OPTIONS else None
            if (
                guias_key not in st.session_state
                or st.session_state.get("tab3_last_case_id") != row.get("ID_Pedido")
            ):
                st.session_state[guias_key] = _guias_val
        st.session_state["tab3_last_case_id"] = row.get("ID_Pedido")

        fecha_recepcion = st.date_input(
            f"📅 Fecha de recepción ({'devolución' if is_dev else 'garantía'})",
            value=st.session_state[fecha_key],
            key=fecha_key,
        )
        estado_opts = ["Sí, completo", "Faltan artículos"]
        estado_recepcion = st.selectbox(
            "📦 ¿Todo llegó correctamente?",
            options=estado_opts,
            index=estado_opts.index(st.session_state[estado_key]) if st.session_state.get(estado_key) in estado_opts else None,
            placeholder="Selecciona el estado de recepción",
            key=estado_key,
        )
        seguimiento_opts = [
            "En revisión de paquete",
            "Falta información vendedor",
            "Autorización de devolución",
            "Pendiente Nota de crédito",
            "Pendiente Transferencia",
            "Cerrado",
        ]

        seguimiento_sel = st.selectbox(
            "🔄 Seguimiento",
            options=seguimiento_opts,
            placeholder="Selecciona el estado de seguimiento",
            key=seg_key,
        )
        if is_dev and guias_key:
            guias_sel = st.selectbox(
                "🚚 Número de guías",
                options=GUIAS_DEVOLUCION_OPTIONS,
                index=(
                    GUIAS_DEVOLUCION_OPTIONS.index(st.session_state[guias_key])
                    if st.session_state.get(guias_key) in GUIAS_DEVOLUCION_OPTIONS
                    else None
                ),
                placeholder="Selecciona cuántas guías se utilizaron",
                key=guias_key,
            )
        else:
            guias_sel = None
        doc_principal = st.file_uploader(
            "🧾 Subir Nota de Crédito / Dictamen (PDF/Imagen)",
            type=["pdf","jpg","jpeg","png"],
            key=f"doc_principal_{row.get('ID_Pedido','')}"
        )
        doc_adicional = st.file_uploader(
            "📂 Subir otro documento (Entrada/Comprobante/Soporte)",
            type=["pdf","jpg","jpeg","png"],
            key=f"doc_extra_{row.get('ID_Pedido','')}"
        )
        comentario_existente = row.get(
            "Comentarios_Admin_Devolucion" if is_dev else "Comentarios_Admin_Garantia",
            ""
        )
        comment_key = f"comentario_admin_{row.get('ID_Pedido','')}"
        if comment_key not in st.session_state:
            st.session_state[comment_key] = comentario_existente
        comentario_admin = st.text_area(
            "📝 Comentario administrativo final",
            key=comment_key,
            value=st.session_state[comment_key]
        )
        submitted = st.form_submit_button("💾 Guardar Confirmación", use_container_width=True)

    # Helper para construir updates
    def update_gsheet_cell(headers, row_idx, col_name, value):
        try:
            col_idx = headers.index(col_name) + 1
        except ValueError:
            return None
        cell = rowcol_to_a1(row_idx, col_idx)
        return {"range": cell, "values": [[value]]}

    # Guardado
    if submitted:
        if not estado_recepcion:
            tab3_alert.warning("⚠️ Completa el campo de estado de recepción.")
            st.stop()
        if not seguimiento_sel:
            tab3_alert.warning("⚠️ Selecciona el estado de seguimiento.")
            st.stop()
        if is_dev:
            guias_val = st.session_state.get(guias_key) if guias_key else None
            if guias_val not in GUIAS_DEVOLUCION_OPTIONS:
                tab3_alert.warning("⚠️ Selecciona la cantidad de guías utilizadas.")
                st.stop()
        else:
            guias_val = None

        # Subir archivos
        tipo_slug = "devolucion" if is_dev else "garantia"
        urls = {}
        carpeta = str(row['ID_Pedido']).strip() or f"caso_{(row.get('Folio_Factura') or 'sfolio')}_{(row.get('Cliente') or 'scliente')}".replace(" ","_")
        for label, file in [("principal", doc_principal), ("extra", doc_adicional)]:
            if file:
                ext = os.path.splitext(file.name)[-1]
                s3_key = f"{carpeta}/{label}_{tipo_slug}_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:4]}{ext}"
                ok, url = upload_file_to_s3(s3_client, S3_BUCKET_NAME, file, s3_key)
                if ok:
                    urls[label] = url

        estado_recepcion_final = "Todo correcto" if estado_recepcion == "Sí, completo" else "Faltan artículos"

        # Columnas por tipo
        if is_dev:
            headers_casos = ensure_sheet_column(
                worksheet_casos,
                headers_casos,
                GUIAS_DEVOLUCION_COL,
            )
        col_fecha_recepcion = pick_first_col(
            headers_casos,
            ["Fecha_Recepcion_Garantia","Fecha_Recepcion_Devolucion"],
            default_if_missing=("Fecha_Recepcion_Devolucion" if is_dev else "Fecha_Recepcion_Garantia")
        )
        col_comentarios = pick_first_col(
            headers_casos,
            ["Comentarios_Admin_Garantia","Comentarios_Admin_Devolucion"],
            default_if_missing=("Comentarios_Admin_Devolucion" if is_dev else "Comentarios_Admin_Garantia")
        )
        col_doc_principal = pick_first_col(
            headers_casos,
            ["Dictamen_Garantia_URL","Nota_Credito_URL"],
            default_if_missing=("Nota_Credito_URL" if is_dev else "Dictamen_Garantia_URL")
        )
        col_doc_extra = pick_first_col(
            headers_casos,
            ["Documento_Adicional_URL"],
            default_if_missing="Documento_Adicional_URL"
        )

        updates = {
            col_fecha_recepcion: fecha_recepcion.strftime("%Y-%m-%d"),
            "Estado_Recepcion": estado_recepcion_final,
            col_doc_principal: urls.get("principal",""),
            col_doc_extra: urls.get("extra",""),
            col_comentarios: comentario_admin,
            "Estado_Caso": "Aprobado",
            "Seguimiento": seguimiento_sel,
        }
        if is_dev and guias_val:
            col_guias = pick_first_col(
                headers_casos,
                [GUIAS_DEVOLUCION_COL],
                default_if_missing=GUIAS_DEVOLUCION_COL,
            )
            updates[col_guias] = guias_val

        ok_all = True
        with st.spinner("Guardando confirmación..."):
            requests = []
            for col, val in updates.items():
                upd = update_gsheet_cell(headers_casos, gsheet_row_idx, col, val)
                if upd:
                    requests.append(upd)
            try:
                if requests:
                    safe_batch_update(worksheet_casos, requests)
            except Exception as e:
                tab3_alert.error(f"❌ Error al actualizar: {e}")
                ok_all = False

        # =================== NUEVO: Generar Word y mostrar descarga (solo Devolución) ===================
        if ok_all and is_dev:
            tab3_alert.success("✅ Confirmación guardada.")
            st.toast("Confirmación guardada", icon="✅")

            try:
                with st.spinner("Generando formato de devolución..."):
                    # Cargar plantilla desde el repo (GitHub → Streamlit)
                    template_path = "plantillas/Formato_Devolución-M.docx"
                    if not os.path.exists(template_path):
                        raise FileNotFoundError(f"No se encontró la plantilla en: {template_path}")

                    doc = Document(template_path)

                    # Mapping exacto a placeholders del .docx
                    mapping = {
                        "Material_Devuelto": _safe_value(row.get("Material_Devuelto")),
                        "Cliente": _safe_value(row.get("Cliente")),
                        "Vendedor_Registro": _safe_value(row.get("Vendedor_Registro")),
                        "Folio_Factura": _safe_value(row.get("Folio_Factura")),
                        "Monto_Devuelto": _safe_value(row.get("Monto_Devuelto")),
                        "Fecha_Recepcion_Devolucion": fecha_recepcion.strftime("%Y-%m-%d"),
                        "Numero_Guias_Devolucion": _safe_value(guias_val),
                        "Area_Responsable": _safe_value(row.get("Area_Responsable")),
                        "Seguimiento": _safe_value(seguimiento_sel),
                        # Para el formato: usamos el URL de la "nota" recién subida como principal (si existe)
                        "Nota_Credito_URL": _safe_value(urls.get("principal", row.get("Nota_Credito_URL",""))),
                        "Folio_Factura_Error": _safe_value(row.get("Folio_Factura_Error")),
                        "Motivo_Detallado": _safe_value(row.get("Motivo_Detallado")),
                        "Comentarios_Admin_Devolucion": _safe_value(comentario_admin),
                    }

                    total_found, total_replaced, remaining_placeholders = _docx_replace_all(doc, mapping)

                    # Guardar en memoria para descarga
                    out_buffer = BytesIO()
                    doc.save(out_buffer)
                    out_buffer.seek(0)

                    st.download_button(
                        label="📄 Descargar Formato de Devolución",
                        data=out_buffer,
                        file_name=f"Formato_Devolucion_{row.get('ID_Pedido','')}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                    pendientes = len(remaining_placeholders)
                    log_message = (
                        f"Placeholders encontrados: {total_found}. "
                        f"Reemplazados: {total_replaced}. "
                        f"Pendientes sin cambio: {pendientes}."
                    )
                    if pendientes:
                        st.warning(
                            log_message
                            + " Sin reemplazo: "
                            + ", ".join(sorted(set(remaining_placeholders)))
                        )
                    else:
                        st.info(log_message)
                    st.info("Puedes completar manualmente 'Imagen producto' y 'Recibe' en el documento descargado.")
            except Exception as e:
                st.error(f"❌ Error al generar el documento: {e}")

            # Importante: NO forzar rerun aquí para que el usuario alcance a descargar.
            st.stop()
        # ================================================================================================

        if ok_all:
            tab3_alert.success("✅ Confirmación guardada.")
            st.session_state["tab3_reload_nonce"] += 1
            get_raw_sheet_data_cached.clear()
            st.toast("Confirmación guardada", icon="✅")
            rerun_current_tab()
        else:
            tab3_alert.error("❌ Ocurrió un problema al guardar.")



# --- TAB 4: CASOS ESPECIALES (Descarga Devoluciones/Garantías) ---
with tab4:
    if tab4_is_active:
        st.session_state[TAB_SESSION_KEY] = 3
        st.session_state["current_tab"] = "3"
    st.header("📥 Casos Especiales (Devoluciones/Garantías)")

    from io import BytesIO
    from datetime import datetime
    import gspread, json, re, math
    import pandas as pd

    # estado local (nonce)
    if "tab4_reload_nonce" not in st.session_state:
        st.session_state["tab4_reload_nonce"] = 0

    # ✅ lector robusto con caché
    @st.cache_data(show_spinner=False, ttl=300)
    def cargar_casos_especiales_cached(sheet_id: str, ws_name: str, _nonce: int):
        ws = safe_open_worksheet(sheet_id, ws_name)
        vals = ws.get_all_values(value_render_option="UNFORMATTED_VALUE")
        if not vals:
            return pd.DataFrame(), [], None
        headers = vals[0]
        df = pd.DataFrame(vals[1:], columns=headers)
        df = df.dropna(how="all")
        for c in ["ID_Pedido", "Cliente", "Folio_Factura", "Tipo_Envio", "Hora_Registro"]:
            if c not in df.columns:
                df[c] = ""
        return df, headers, ws

    # 🔁 recargar
    col_a, col_b = st.columns([1, 5])
    with col_a:
        if st.button("🔄 Recargar Casos", type="secondary", key="tab4_reload_btn"):
            if allow_refresh("tab4_last_refresh"):
                st.session_state["tab4_reload_nonce"] += 1
                st.toast("♻️ Casos recargados.", icon="♻️")
                rerun_current_tab()

    # leer hoja
    prev_nonce = st.session_state["tab4_reload_nonce"]
    try:
        df_ce, headers_ce, ws_casos = cargar_casos_especiales_cached(
            GOOGLE_SHEET_ID, "casos_especiales", prev_nonce
        )
        st.session_state["_lastgood_casos_especiales"] = (
            df_ce.copy(), headers_ce
        )
    except gspread.exceptions.WorksheetNotFound:
        st.error("❌ No existe la hoja 'casos_especiales'.")
        df_ce, headers_ce, ws_casos = pd.DataFrame(), [], None
    except gspread.exceptions.APIError as e:
        st.session_state["tab4_reload_nonce"] = max(0, prev_nonce - 1)
        snap = st.session_state.get("_lastgood_casos_especiales")
        if snap:
            st.warning(
                "♻️ Google Sheets dio un error temporal al leer 'casos_especiales'. Mostrando el último dato bueno en caché."
            )
            df_ce, headers_ce = snap
            ws_casos = None
        else:
            st.error(f"❌ Error al leer 'casos_especiales': {e}")
            df_ce, headers_ce, ws_casos = pd.DataFrame(), [], None

    if df_ce.empty:
        st.info("ℹ️ No hay registros en 'casos_especiales'.")
        st.stop()

    # ------- Normalizador de URLs (Adjuntos puede venir como JSON/CSV/texto) -------
    def _normalize_urls(value):
        if value is None:
            return []
        if isinstance(value, float) and math.isnan(value):
            return []
        s = str(value).strip()
        if not s or s.lower() in ("nan", "none", "n/a"):
            return []
        urls = []
        try:
            obj = json.loads(s)
            if isinstance(obj, list):
                for it in obj:
                    if isinstance(it, str) and it.strip():
                        urls.append(it.strip())
                    elif isinstance(it, dict):
                        u = it.get("url") or it.get("URL") or it.get("href") or it.get("link")
                        if u and str(u).strip():
                            urls.append(str(u).strip())
            elif isinstance(obj, dict):
                for k in ("url", "URL", "link", "href"):
                    if obj.get(k):
                        urls.append(str(obj[k]).strip())
        except Exception:
            parts = re.split(r"[,\n;]+", s)
            for p in parts:
                p = p.strip()
                if p:
                    urls.append(p)
        # únicos
        seen, out = set(), []
        for u in urls:
            if u not in seen:
                seen.add(u)
                out.append(u)
        return out

    # ------- Derivados de enlaces y campos mínimos -------
    for c in [
        "Adjuntos", "Hoja_Ruta_Mensajero", "Nota_Credito_URL",
        "Documento_Adicional_URL", "Dictamen_Garantia_URL",
        "Estado", "Estado_Caso", "Estado_Recepcion", "Tipo_Envio_Original", "Estatus_OrigenF",
        "Resultado_Esperado", "Material_Devuelto", "Monto_Devuelto", "Motivo_Detallado",
        "Numero_Serie", "Fecha_Compra", "Numero_Cliente_RFC", "Area_Responsable", "Nombre_Responsable", "Turno", "Fecha_Entrega",
        "Seguimiento",
        GUIAS_DEVOLUCION_COL
    ]:
        if c not in df_ce.columns:
            df_ce[c] = ""

    # Links listos para tabla/Excel
    df_ce["Links_Adjuntos"] = df_ce["Adjuntos"].apply(lambda v: "\n".join(_normalize_urls(v)) if str(v).strip() else "")
    df_ce["Link_Guia"] = df_ce["Hoja_Ruta_Mensajero"].astype(str).fillna("")
    # prioriza dictamen garantía; si no, nota crédito
    df_ce["Link_Dictamen_o_Nota"] = df_ce.apply(
        lambda r: (str(r.get("Dictamen_Garantia_URL","")).strip() or str(r.get("Nota_Credito_URL","")).strip()),
        axis=1
    )
    df_ce["Link_Doc_Adicional"] = df_ce["Documento_Adicional_URL"].astype(str).fillna("")

    # ------- Filtros rápidos -------
    colf1, colf2, colf3 = st.columns([1.2, 1.4, 2.4])
    with colf1:
        filtro_tipo = st.selectbox(
            "Tipo de caso",
            options=["Todos", "🔁 Devolución", "🛠 Garantía"],
            index=0
        )
    seguimiento_options = sorted(
        {
            str(val).strip()
            for val in df_ce.get("Seguimiento", [])
            if str(val).strip() and str(val).strip().lower() not in {"nan", "none"}
        }
    )
    seguimiento_options = ["Todos"] + seguimiento_options
    with colf2:
        filtro_seguimiento = st.selectbox(
            "Seguimiento",
            options=seguimiento_options,
            index=0,
        )
    with colf3:
        term = st.text_input("Buscar (Cliente / Folio )", "")

    df_view = df_ce.copy()

    if filtro_tipo != "Todos":
        # soporta tanto Tipo_Envio como Tipo_Caso
        tipo_col = "Tipo_Envio" if "Tipo_Envio" in df_view.columns else ("Tipo_Caso" if "Tipo_Caso" in df_view.columns else None)
        if tipo_col:
            df_view = df_view[df_view[tipo_col].astype(str).str.strip() == filtro_tipo]

    if "Seguimiento" in df_view.columns and filtro_seguimiento != "Todos":
        df_view = df_view[
            df_view["Seguimiento"].astype(str).str.strip() == filtro_seguimiento
        ]

    if term.strip():
        t = term.strip().lower()
        df_view = df_view[
            df_view["Cliente"].astype(str).str.lower().str.contains(t, na=False) |
            df_view["Folio_Factura"].astype(str).str.lower().str.contains(t, na=False) |
            df_view["ID_Pedido"].astype(str).str.lower().str.contains(t, na=False)
        ]

    # ------- Ordenar: últimos primero por Hora_Registro si existe -------
    def _to_dt(s):
        return pd.to_datetime(s, errors='coerce', dayfirst=True, infer_datetime_format=True)

    if "Hora_Registro" in df_view.columns:
        dth = _to_dt(df_view["Hora_Registro"])
        if dth.notna().any():
            df_view = df_view.assign(_dt=dth).sort_values("_dt", ascending=False, na_position='last').drop(columns="_dt")
        else:
            df_view = df_view.iloc[::-1].reset_index(drop=True)
    else:
        df_view = df_view.iloc[::-1].reset_index(drop=True)

    # ------- Métricas rápidas -------
    c1, c2, c3, c4 = st.columns(4)
    with c1: st.metric("Casos totales", len(df_ce))
    with c2: st.metric("Devoluciones", len(df_ce[df_ce["Tipo_Envio"].astype(str).str.contains("Devoluci", case=False, na=False)]))
    with c3: st.metric("Garantías", len(df_ce[df_ce["Tipo_Envio"].astype(str).str.contains("Garant", case=False, na=False)]))
    with c4: st.metric("Completados", len(df_ce[df_ce["Estado"].astype(str).str.strip() == "🟢 Completado"]))

    st.markdown("---")

    # ------- Columnas a mostrar/descargar -------
    columnas_base = [
        "ID_Pedido","Hora_Registro","Vendedor_Registro","Cliente","Folio_Factura",
        "Numero_Serie","Fecha_Compra",
        "Tipo_Envio","Estado","Estado_Caso","Estado_Recepcion","Seguimiento",
        "Tipo_Envio_Original","Estatus_OrigenF","Turno","Fecha_Entrega",
        "Resultado_Esperado","Material_Devuelto","Monto_Devuelto","Motivo_Detallado",
        "Numero_Cliente_RFC","Area_Responsable","Nombre_Responsable",
        GUIAS_DEVOLUCION_COL
    ]
    columnas_links = ["Links_Adjuntos","Link_Guia","Link_Dictamen_o_Nota","Link_Doc_Adicional"]

    columnas_existentes = [c for c in columnas_base + columnas_links if c in df_view.columns]
    if not columnas_existentes:
        columnas_existentes = df_view.columns.tolist()

    st.success(f"✅ {len(df_view)} casos listados (últimos primero).")

    st.dataframe(
        df_view[columnas_existentes],
        use_container_width=True, hide_index=True
    )

    # ------- Descargar Excel -------
    output_casos = BytesIO()
    with pd.ExcelWriter(output_casos, engine="xlsxwriter") as writer:
        df_view[columnas_existentes].to_excel(writer, index=False, sheet_name="casos_especiales")
        # (opcional) se podría formatear celdas/hipervínculos aquí
    data_xlsx = output_casos.getvalue()

    st.download_button(
        label="📥 Descargar Excel Casos Especiales (últimos primero)",
        data=data_xlsx,
        file_name=f"casos_especiales_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
